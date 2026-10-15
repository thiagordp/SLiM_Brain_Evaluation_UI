"""python-docx helpers shared by the guide and the annex.

Restrained academic style: Calibri, navy headings (as the Training report),
hairline horizontal rules in tables, status chips that always carry their text
label (colour is never the only signal). Contrast of each chip is checked by
``contrast_report``.
"""
from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
NAVY, BLUE, INK, MUTED, RULE, PANEL = "1F3864", "2E5395", "1F2328", "59606A", "C9CED6", "F3F5F8"
STATUS_STYLE = {   # text colour, fill
    "AGREED": ("1D5B36", "DCEFE2"),
    "PROVISIONAL": ("6E4300", "FBE9C6"),
    "OPEN": ("8A1C2B", "F8DDE1"),
}
STATUS_LABEL = {"AGREED": "Agreed", "PROVISIONAL": "Provisional", "OPEN": "Open"}
IMPL_LABEL = {"Implemented": "In the app", "Reported implemented": "Reported", "Not implemented": "Not in the app",
              "Unknown": "Unknown", "n/a": "Principle"}


# ------------------------------------------------------------------ document
def new_document(landscape: bool = False, margins_cm: float = 1.8) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(margins_cm))
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.6)
    sec.header_distance, sec.footer_distance = Cm(0.9), Cm(0.8)

    st = doc.styles
    normal = st["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    _east_asia(normal.element, FONT)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(3), 1.08
    for name, size, color, before, after in (("Title", 22, NAVY, 0, 2), ("Heading 1", 14.5, NAVY, 10, 4),
                                             ("Heading 2", 11.5, BLUE, 8, 3), ("Heading 3", 10, BLUE, 6, 2)):
        s = st[name]
        s.font.name, s.font.size, s.font.bold = FONT, Pt(size), name != "Title"
        s.font.color.rgb = RGBColor.from_string(color)
        s.font.italic = False
        _east_asia(s.element, FONT)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True
        if name == "Title":
            # The built-in Title style carries a bottom border; keep the page clean.
            ppr = s.element.get_or_add_pPr()
            for b in ppr.findall(qn("w:pBdr")):
                ppr.remove(b)
    return doc


def _east_asia(style_el, font):
    rpr = style_el.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), font)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]


def header_footer(doc: Document, left: str, right: str, footer_left: str) -> None:
    for sec in doc.sections:
        hp = sec.header.paragraphs[0]
        hp.text = ""
        hp.style = doc.styles["Normal"]
        _tab_line(hp, left, right, sec, size=8, color=MUTED, rule="bottom")
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        fp.style = doc.styles["Normal"]
        width = sec.page_width - sec.left_margin - sec.right_margin
        fp.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)
        r = fp.add_run(footer_left)
        style_run(r, size=8, color=MUTED)
        r = fp.add_run("\tPage ")
        style_run(r, size=8, color=MUTED)
        field(fp, "PAGE", size=8, color=MUTED)
        r = fp.add_run(" of ")
        style_run(r, size=8, color=MUTED)
        field(fp, "NUMPAGES", size=8, color=MUTED)


def _tab_line(p, left, right, sec, size, color, rule=None):
    width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)
    style_run(p.add_run(left), size=size, color=color, bold=True)
    style_run(p.add_run("\t" + right), size=size, color=color)
    if rule:
        para_border(p, rule, color=RULE, sz=4, space=3)


def field(p, code: str, size=8, color=MUTED):
    def r_el():
        r = p.add_run()
        style_run(r, size=size, color=color)
        return r
    r = r_el()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = r_el()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {code} "; r._r.append(it)
    r = r_el()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    r = r_el(); r.text = "1"
    r = r_el()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)


# ------------------------------------------------------------------ runs and paragraphs
def style_run(r, size=None, color=None, bold=None, italic=None, caps=None, font=FONT):
    r.font.name = font
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), font)
    if size is not None:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        r.font.bold = bold
    if italic is not None:
        r.font.italic = italic
    if caps is not None:
        r.font.small_caps = caps
    return r


def run_shading(r, fill):
    rpr = r._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def run_border(r, color, sz=4):
    rpr = r._r.get_or_add_rPr()
    b = OxmlElement("w:bdr")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), color)
    rpr.append(b)


def para(container, text="", size=None, color=None, bold=None, italic=None, after=None, before=None,
         align=None, keep=False, indent_cm=None, style=None):
    p = container.add_paragraph(style=style) if style else container.add_paragraph()
    if text:
        rich(p, text, size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    if after is not None:
        pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if keep:
        pf.keep_with_next = True
    if indent_cm is not None:
        pf.left_indent = Cm(indent_cm)
    return p


def rich(p, text, size=None, color=None, bold=None, italic=None):
    """Append text with **bold** markup (no underscore markup: ids and file names contain underscores)."""
    import re
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            style_run(p.add_run(part[2:-2]), size=size, color=color, bold=True, italic=italic)
        else:
            style_run(p.add_run(part), size=size, color=color, bold=bold, italic=italic)
    return p


def label_line(container, label, text, size=9, label_color=BLUE, after=1.5):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    style_run(p.add_run(label + "  "), size=size - 0.5, color=label_color, bold=True)
    rich(p, text, size=size)
    return p


def chip(p, status, size=7.5):
    fg, bg = STATUS_STYLE[status]
    r = p.add_run(" " + STATUS_LABEL[status] + " ")
    style_run(r, size=size, color=fg, bold=True)
    run_shading(r, bg)
    return r


def impl_tag(p, impl, size=7.5):
    r = p.add_run(IMPL_LABEL[impl])
    style_run(r, size=size, color=MUTED, italic=True)
    return r


def tag(p, pid, anchor=None, size=7):
    """A small grey rule id, linked to its entry when an anchor is given."""
    if anchor:
        return link(p, pid, anchor, size=size, color=MUTED)
    return style_run(p.add_run(pid), size=size, color=MUTED)


def para_border(p, side, color=RULE, sz=4, space=1):
    ppr = p._p.get_or_add_pPr()
    pb = ppr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr"); ppr.append(pb)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), str(space)); b.set(qn("w:color"), color)
    pb.append(b)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ bookmarks and links
_BM = [100]


def bookmark(p, name):
    _BM[0] += 1
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(_BM[0])); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(_BM[0]))
    p._p.insert(1 if p._p.pPr is not None else 0, s)
    p._p.append(e)


def link(p, text, anchor, size=None, color=BLUE, bold=None):
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    h.set(qn("w:history"), "1")
    r = p.add_run(text)
    style_run(r, size=size, color=color, bold=bold)
    h.append(r._r)
    p._p.append(h)
    return h


def bm_name(pid: str) -> str:
    return "pol_" + pid.replace("-", "_")


# ------------------------------------------------------------------ tables
def table(container, widths_cm, header=None, header_fill=PANEL, size=8.5, rules=True, outer=True, pad=45):
    """A fixed-layout table with hairline horizontal rules. Returns the Table."""
    ncols = len(widths_cm)
    t = container.add_table(rows=0, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tbl = t._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths_cm) * 567))); tw.set(qn("w:type"), "dxa")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(tw)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "insideH", "left", "right", "insideV"):
        b = OxmlElement(f"w:{side}")
        if (side in ("top", "bottom") and outer) or (side == "insideH" and rules):
            b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), RULE)
        else:
            b.set(qn("w:val"), "nil")
        borders.append(b)
    tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", pad), ("bottom", pad), ("left", 85), ("right", 85)):
        m = OxmlElement(f"w:{side}"); m.set(qn("w:w"), str(w)); m.set(qn("w:type"), "dxa"); mar.append(m)
    tblPr.append(mar)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    t._widths = widths_cm
    t._size = size
    if header:
        row = add_row(t, header, header=True, fill=header_fill, bold=True, color=NAVY, size=size - 0.5)
    return t


def add_row(t, cells, header=False, fill=None, bold=None, color=None, size=None, cant_split=True, keep=False):
    """cells: list of str or callables(cell) that fill the cell themselves."""
    row = t.add_row()
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader"); trPr.append(th)
    size = size or t._size
    for i, (cell, w) in enumerate(zip(row.cells, t._widths)):
        cell.width = Cm(w)
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
        tcW.set(qn("w:w"), str(int(w * 567))); tcW.set(qn("w:type"), "dxa")
        if fill:
            shade(cell, fill)
        content = cells[i] if i < len(cells) else ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if keep:
            p.paragraph_format.keep_with_next = True
        if callable(content):
            content(cell)
        else:
            lines = str(content).split("\n")
            for k, line in enumerate(lines):
                q = p if k == 0 else cell.add_paragraph()
                q.paragraph_format.space_after = Pt(0)
                rich(q, line, size=size, color=color, bold=bold)
    return row


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_border(cell, side, color, sz=12):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders"); tcPr.append(tb)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:color"), color); b.set(qn("w:space"), "0")
    tb.append(b)


def merge_row(row):
    a = row.cells[0]
    return a.merge(row.cells[-1])


def keep_table_together(t):
    """Keep every row but the last with the next one (short tables only)."""
    rows = t.rows
    for row in rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


# ------------------------------------------------------------------ accessibility
def _lum(hexcolor):
    c = [int(hexcolor[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(fg, bg):
    a, b = sorted((_lum(fg), _lum(bg)), reverse=True)
    return (a + 0.05) / (b + 0.05)


def contrast_report():
    rows = [(k, contrast(fg, bg)) for k, (fg, bg) in STATUS_STYLE.items()]
    rows += [("muted text on white", contrast(MUTED, "FFFFFF")), ("blue on white", contrast(BLUE, "FFFFFF")),
             ("ink on panel", contrast(INK, PANEL))]
    return rows
