"""Render the annex: the full calibration decision register and its evidence.

    python app/tools/calibration_guide/build_annex.py BUILD_DIR OUT.docx
"""
from __future__ import annotations

import json
import pathlib
import re
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import docstyle as ds  # noqa: E402
import policies as pol  # noqa: E402
from docx.shared import Pt  # noqa: E402
from report_docx import EVALUATORS  # noqa: E402

TITLE = "SLiM Brain Calibration Register"
W = 26.7                                     # A4 landscape text width at 1.5 cm margins (cm)
ABBR = dict(zip(EVALUATORS, ["Fr", "Tb", "Tg", "Al", "Gp", "Gv", "Vj"]))
SHORT_ANS = {"Yes": "Y", "No": "N", "In part": "P", "": "–", "none": "–"}
IDS = re.compile(r"\b((?:GEN|CL|GR|FD|CO|RE|RC|CC)-\d+)\b")
POL = pol.by_id()
CHK = {c["id"] for c in pol.CONSISTENCY_CHECKS}
CATEGORY_TEXT = [
    ("Yes, No or In part", "CONFIRMED", "A plain answer, sometimes followed by a note. It is compared with the evaluators' answers."),
    ("No, with a proposed value", "REPLACEMENT", "The Final gives another value for the field, such as Other for an assigned law, "
                                                  "or “No. Abstract”. The assigned value is judged No, and the proposed value is shown "
                                                  "separately."),
    ("Unresolved", "UNRESOLVED", "Marked “??”, “No?” or “To be discussed”, cut off, or contradicted by a comment on the same "
                                 "cell. Treated as open and never compared."),
    ("Not applicable", "NOT_APPLICABLE", "“N/A”. The Final judges Relation type No, so Direction is not asked."),
    ("No Final", "EMPTY", "Nothing was entered."),
    ("Missing Concepts", "SET_VALUED", "Q14 proposals. Annex C gives one row per proposed Concept."),
    ("Paper level", "PAPER_LEVEL", "Claim recall, Restatements and Missing Claims. See annex D."),
]
STATS_NOTES = [
    "Skipped answers are excluded, so the number of rated pairs differs between questions. Report it with each coefficient.",
    "Most Finals are proposals by one reviewer. Agreement with a Final measures convergence, not correctness.",
    "A majority among seven evaluators is not a calibration decision. This register never uses one to settle an open item.",
    "Answers revised after reading the Finals are no longer independent. Agreement before and after calibration should be "
    "reported separately.",
    "Thibault noted on 11 October that the results describe the current models and may not hold for later ones. This "
    "belongs in the Discussion.",
]


def text(p, s, size=8, color=None, bold=None):
    for part in IDS.split(s):
        if not part:
            continue
        if IDS.fullmatch(part) and part in POL:
            ds.link(p, part, ds.bm_name(part), size=size, color=ds.BLUE, bold=bold)
        elif IDS.fullmatch(part) and part in CHK:
            ds.link(p, part, "chk_" + part.replace("-", "_"), size=size, color=ds.BLUE, bold=bold)
        else:
            ds.rich(p, part, size=size, color=color, bold=bold)


def h1(doc, s, brk=True):
    if brk:
        ds.page_break(doc)
    return doc.add_paragraph(s, style="Heading 1")


def note(doc, s, size=8.6, after=3, color=ds.MUTED):
    p = ds.para(doc, "", after=after)
    text(p, s, size=size, color=color)
    return p


def bullet(doc, s, size=8.6):
    q = doc.add_paragraph(style="List Bullet")
    q.paragraph_format.space_after = Pt(1.5)
    text(q, s, size=size)


def read_as(r):
    c = r["final_category"]
    if c == "CONFIRMED":
        s = r["final_value"]
    elif c == "REPLACEMENT":
        s = f'No; proposes {r["final_replacement"]}'
    elif c == "UNRESOLVED":
        s = "Unresolved" + (f' (tentative {r["final_value"]})' if r["final_value"] else "")
    elif c == "NOT_APPLICABLE":
        s = "Not applicable"
    elif c == "EMPTY":
        s = "No Final"
    elif c == "SET_VALUED":
        s = "See annex C"
    elif c == "PAPER_LEVEL":
        s = "See annex D"
    else:
        s = c.replace("_", " ").capitalize()
    flags = [f for f in r["final_flags"] if f not in ("TENTATIVE",)]
    if flags:
        s += "\n" + ", ".join(f.lower().replace("_", " ") for f in flags)
    return s


def answers_cell(r, e):
    a = r["answers"][e]
    if r["qid"] == "Q14":
        a = r["evaluators"][e]["answer"]
        return "–" if a.lower() in ("none", "") else "+"
    s = SHORT_ANS.get(a, a[:6])
    return s + ("†" if e in r["gated"] else "")


def build(reg, out):
    doc = ds.new_document(landscape=True, margins_cm=1.5)
    ds.header_footer(doc, "SLiM Brain · Calibration Register (annex to the Calibration Guide)",
                     "Instrument 4.1 · Training phase", "Sources up to 14 October 2026")
    rows = reg["rows"]
    items = [r for r in rows if not r["paper_level"]]

    doc.add_paragraph(TITLE, style="Title")
    ds.para(doc, "Annex to the SLiM Brain Calibration Guide · human evaluation, Training phase · instrument 4.1 · "
                 "papers SRC-0006 and SRC-0009 · sources up to 14 October 2026", size=10, color=ds.MUTED, after=6)
    p = ds.para(doc, "", after=4)
    text(p, f"This annex holds the detail behind the guide. It lists every contested item in Appendix B of the Training "
            f"report ({len(rows)} rows: {len(items)} items on Claims and relations, and {len(rows) - len(items)} on whole "
            f"papers), with the seven answers, the Final as written and how it was read. It then gives each evaluator's "
            f"review list, the contradictions found in the Finals, all reviewer comments, and the evidence for every rule.",
         size=9.5)
    p = ds.para(doc, "", after=6)
    text(p, "The register was generated from Calibration/TRAINING_AGREEMENT_REPORT.docx as it stood on 14 October 2026, "
            "by the scripts in app/tools/calibration_guide. The source files were not changed.", size=9, color=ds.MUTED)

    # ---------------- A
    doc.add_paragraph("A  How to read this register", style="Heading 1")
    doc.add_paragraph("Status of each item", style="Heading 2")
    t = ds.table(doc, [3.0, W - 3.0], size=8.6)
    for st in (pol.AGREED, pol.PROVISIONAL, pol.OPEN):
        ds.add_row(t, [lambda c, st=st: ds.chip(c.paragraphs[0], st, size=7.5),
                       {"AGREED": "A meeting recorded the outcome, with at least two participants agreeing. This applies to four "
                                  "items: CLM-0006-001 (Pseudolaw and Legal literacy), CLM-0006-002 (Direction) and "
                                  "CLM-0006-015 (textual grounding).",
                        "PROVISIONAL": "The Final is one reviewer's proposal. Vaclav proposed the SRC-0006 Finals in his email of "
                                       "9 October. The author of the SRC-0009 Finals is not recorded, though the reviewer "
                                       "comments there are Thiago's. A proposal is not a team decision.",
                        "OPEN": "The Final is tentative, missing, cut off or self-contradictory, or a reviewer marked its cell "
                                "yellow. Vaclav wrote on 9 October: “There is few open question for me (highlighted in "
                                "yellow)”. The proposed value is kept but not used for comparison."}[st]])
    ds.keep_table_together(t)

    doc.add_paragraph("How each Final was read", style="Heading 2")
    t = ds.table(doc, [4.2, W - 4.2], header=["Read as", "Meaning"], size=8.6)
    for label, _key, v in CATEGORY_TEXT:
        ds.add_row(t, [label, v])
    ds.keep_table_together(t)

    doc.add_paragraph("Columns, symbols and colours", style="Heading 2")
    for s in [
        "**Answers.** Y = Yes, N = No, P = In part, – = none. For Missing Concepts, + means the evaluator proposed at least "
        "one Concept. Evaluators: " + ", ".join(f"{v} {k}" for k, v in ABBR.items()) + ".",
        "**Skipped answers (†).** The Training answers were stored before the 4.1 gates existed. The register reads them as "
        "the app does. An evaluator who answered No to Q02 has no applicable answer for Q03–Q11 of that Claim, and one who "
        "answered No to Relation type has none for its Grounding and Direction. These answers carry † and are never "
        "compared with a Final. When the gate question is not listed because all seven evaluators agreed, nothing is "
        "skipped.",
        "**Differs from Final.** The evaluators whose applicable answer differs from the Final. It is filled only when the "
        "item is not open and the Final reads as Yes, No or In part, or as No with a proposed value. Otherwise the cell "
        "gives the reason it is empty. validate.py recomputes it from the seven answers.",
        "**Yellow.** The report generator shades ties pale yellow (FFF2CC) and minority answers orange (FCE4D6). It never "
        "produces bright yellow (FFFF00) or highlighting, so those are reviewers' marks.",
        "**Rules.** The ids of the guide rules that apply to the item. They link to annex I, and CC ids link to the "
        "contradictions in annex G.",
    ]:
        bullet(doc, s)

    doc.add_paragraph("Sources", style="Heading 2")
    t = ds.table(doc, [2.0, W - 2.0], size=8.2)
    for k, v in pol.SOURCES.items():
        ds.add_row(t, [lambda c, k=k: ds.style_run(c.paragraphs[0].add_run(k), size=8.2, color=ds.NAVY, bold=True), v])

    # ---------------- B
    h1(doc, "B  Item register")
    note(doc, "One block per Claim, in the order of the Training report. Relation rows keep the report's order, with "
              "Direction before Relation type. Finals are quoted as written, and a line break inside a Final is shown as "
              "“ / ”.")
    widths = [0.8, 2.6, 3.0] + [1.05] * 7 + [5.2, 2.4, 1.9, 2.3, 1.15]
    head = ["#", "Question", "Assigned"] + [ABBR[e] for e in EVALUATORS] + ["Final as written", "Read as", "Status",
                                                                             "Differs from Final", "Rules"]
    for paper in ("SRC-0006", "SRC-0009"):
        doc.add_paragraph(paper, style="Heading 2")
        t = ds.table(doc, widths, header=head, size=7.6)
        current = None
        for r in [x for x in items if x["paper"] == paper]:
            if r["id"] != current:
                current = r["id"]
                st = reg["statements"].get(current, "")
                row = ds.add_row(t, [""] * len(widths), fill=ds.PANEL, keep=True)
                cell = ds.merge_row(row)
                cell.paragraphs[0].text = ""
                for extra in cell.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
                q = cell.paragraphs[0]
                ds.style_run(q.add_run(current + "  "), size=8, color=ds.NAVY, bold=True)
                ds.bookmark(q, "clm_" + current.replace("-", "_"))
                ds.style_run(q.add_run(st), size=7.4, color=ds.INK, italic=True)
            fin = " / ".join(p.strip() for p in r["final_paras"] if p.strip()) or "—"
            diff = ", ".join(ABBR[e] for e in r["differs"]) or ("—" if not r["differs_note"] else "")
            dn = r["differs_note"]
            if dn.startswith("Not compared"):
                diff = dn.replace("Not compared: ", "not compared: ")
            elif dn:
                diff = (diff + "\n" if diff else "") + dn.replace("Gated (not compared): ", "skipped: ")
            cells = [str(r["n"]), r["question"], r["assigned"] or "—"] + [answers_cell(r, e) for e in EVALUATORS] + [
                fin, read_as(r), lambda c, r=r: ds.chip(c.paragraphs[0], r["status"], size=6.6), diff,
                lambda c, r=r: text(c.paragraphs[0], ", ".join(r["policy_ids"][:3]), size=7.2)]
            ds.add_row(t, cells)

    # ---------------- C
    h1(doc, "C  Missing-Concept proposals (Q14)")
    note(doc, "One row per proposed Concept. “Considered reasonable” means the Final called the addition reasonable or "
              "acceptable. It does not mean the Concept was judged necessary or missing, or that it was adopted. Some "
              "report cells had lost their line breaks; they were split against the Concept vocabulary, and every split "
              "matched a known Concept name exactly.")
    disp = {"AGREED_MISSING": "Agreed as missing (meeting)", "NOT_ADOPTED": "Not adopted",
            "NOT_ADOPTED_OPEN_POLICY": "Not adopted; policy question open", "CONSIDERED_REASONABLE": "Considered reasonable",
            "REASONABLE_WITH_CHANGE": "Reasonable, with a change", "UNCLEAR": "Unclear"}
    t = ds.table(doc, [2.4, 2.0, 4.6, 1.4, 8.0, 4.4, 3.9], header=["Claim", "Evaluator", "Concept", "New?",
                                                                 "Final as written", "Outcome", "Basis"], size=7.8)
    for p in reg["proposals"]:
        ds.add_row(t, [p["claim"], p["evaluator"], p["concept"], "yes" if p["new"] else "", p["final"],
                       disp[p["disposition"]], p["basis"]])

    # ---------------- D
    h1(doc, "D  Paper-level items: Claim recall, Restatements, Missing Claims")
    note(doc, "These items concern each paper as a whole. The team has not decided how to judge them, and no agreement is "
              "sought on which Claims are missing. Texts are quoted as written.")
    for paper in ("SRC-0006", "SRC-0009"):
        doc.add_paragraph(paper, style="Heading 2")
        pl = [r for r in rows if r["id"] == paper]
        t = ds.table(doc, [3.4] + [2.6] * 7 + [5.1], header=["Question"] + EVALUATORS + ["Final"], size=7.8)
        for r in pl:
            if r["qid"] == "MISSING_CLAIMS":
                continue
            ds.add_row(t, [r["question"]] + [r["evaluators"][e]["answer"] for e in EVALUATORS] + [r["final"] or "— (marked yellow)" if r["manual_yellow"] and not r["final"] else r["final"]])
        mc = [r for r in pl if r["qid"] == "MISSING_CLAIMS"][0]
        q = ds.para(doc, "", before=4, after=2, keep=True)
        ds.style_run(q.add_run("Missing Claims, as written"), size=9, color=ds.BLUE, bold=True)
        t = ds.table(doc, [3.0, W - 3.0], size=7.8)
        for e in EVALUATORS:
            ev = mc["evaluators"][e]
            body = (ev["answer"] + ("\n" + ev["comment"] if ev["comment"] else "")).strip()
            if ev["answer"] == "not asked":
                continue
            ds.add_row(t, [e, body.replace("listed\n", "")], cant_split=False)
        ds.add_row(t, ["Not asked", ", ".join(e for e in EVALUATORS if mc["evaluators"][e]["answer"] == "not asked")
                       + " (they answered All on the recall scale)"])
        ds.add_row(t, ["Final", mc["final"] or "—"])

    # ---------------- E
    h1(doc, "E  Flagged Finals")
    note(doc, "Finals that are cut off, contradictory, tentative, empty, marked yellow, or corrected for case. Nothing has "
              "been repaired. Each Final is shown as written.")
    t = ds.table(doc, [2.4, 3.2, 4.0, 6.0, 3.2, W - 18.8], header=["Claim", "Question", "Assigned", "Final as written", "Flag", "How it was read"], size=7.8)
    for r in rows:
        flags = list(r["final_flags"])
        if r["final_category"] == "EMPTY":
            flags.append("EMPTY")
        if not flags and r["final_category"] != "UNRESOLVED":
            continue
        reading = r["final_reason"] or r["status_basis"]
        if "YELLOW_EVALUATOR_CELL" in flags:
            reading = ("Only an evaluator cell (" + ", ".join(c for c in r["manual_yellow_cells"]) + ") is marked yellow. "
                       "Nobody recorded what the mark means, so the Final is kept as proposed.")
        if "CASE_NORMALISED" in flags:
            reading = "Lower-case “yes” read as Yes."
        ds.add_row(t, [r["id"], r["question"], r["assigned"] or "—", " / ".join(x for x in r["final_paras"] if x.strip()) or "—",
                       ", ".join(f.lower().replace("_", " ") for f in flags), reading])
    ds.add_row(t, ["CLM-0006-001", "Missing Concepts (Q14)", "selected / proposed", "Legal Literatcy", "typo",
                   "Read as “Legal literacy”, which was agreed at the 7 October meeting (23:03–23:22). The text is kept as written."])

    # ---------------- F
    h1(doc, "F  Review list per evaluator")
    note(doc, "For each evaluator, the items where their applicable answer differs from a Final that is not open. "
              "Reconsider the answer, or add a comment that explains it. Most of these Finals are one reviewer's proposal, "
              "so differing from one is not an error. Items with an open Final are only counted here, because they are "
              "not corrections.")
    open_items = [r for r in items if r["status"] == "OPEN"]
    for e in EVALUATORS:
        mine = [r for r in items if e in r["differs"]]
        doc.add_paragraph(e, style="Heading 2")
        n_open = sum(1 for r in open_items if r["answers"][e])
        n_gated = sum(1 for r in items if e in r["gated"])
        n_agreed = sum(1 for r in mine if r["status"] == "AGREED")
        note(doc, f"{len(mine)} answers differ from a Final that is not open, and {n_agreed} of these Finals were agreed at a "
                  f"meeting. {n_open} items wait on an open decision. {n_gated} stored answers are skipped by the gates.",
             size=8.4, after=2)
        if not mine:
            continue
        t = ds.table(doc, [0.9, 2.4, 3.3, 4.0, 1.6, 7.0, 2.2, W - 21.4], header=["#", "Claim", "Question", "Assigned",
                                                                                "Answer", "Final as written", "Status", "Rules"], size=7.6)
        for r in mine:
            ds.add_row(t, [str(r["n"]), r["id"], r["question"], r["assigned"] or "—", r["answers"][e],
                           " / ".join(x for x in r["final_paras"] if x.strip()),
                           lambda c, r=r: ds.chip(c.paragraphs[0], r["status"], size=6.6),
                           lambda c, r=r: text(c.paragraphs[0], ", ".join(r["policy_ids"][:3]), size=7.2)])

    # ---------------- G
    h1(doc, "G  Contradictions in the Finals")
    note(doc, "Points where a Final conflicts with another Final, with a meeting, with the schema or with the app. None of "
              "them is resolved here. The first five were already raised in the earlier summary of 14 October.")
    t = ds.table(doc, [1.6, 5.0, W - 9.6, 3.0], header=["Check", "Items", "What conflicts, with the evidence", "Rules"], size=7.9)
    for c in pol.CONSISTENCY_CHECKS:
        def cid(cell, c=c):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run(c["id"]), size=8.2, color=ds.NAVY, bold=True)
            ds.bookmark(q, "chk_" + c["id"].replace("-", "_"))
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0)
            ds.style_run(q2.add_run("earlier summary" if c["origin"].startswith("SUM") else "new"), size=6.8, color=ds.MUTED)

        def body(cell, c=c):
            text(cell.paragraphs[0], c["point"], size=7.9)
            for k, s in c["evidence"]:
                q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(0)
                ds.style_run(q.add_run(f"{k}: “{s}”"), size=7.2, color=ds.MUTED)
        ds.add_row(t, [cid, ", ".join(c["items"]) or "—", body,
                       lambda cell, c=c: text(cell.paragraphs[0], ", ".join(c["policies"]), size=7.6)])
    doc.add_paragraph("Implications for the agreement statistics", style="Heading 2")
    note(doc, "These are observations made while building the register, not team decisions.", after=2)
    for s in STATS_NOTES:
        bullet(doc, s)

    # ---------------- H
    h1(doc, "H  Reviewer comments in the Training report")
    note(doc, f'All {len(reg["comments"])} Word comments, quoted as written, with the cell each one is attached to. The ids '
              f"are the report's own comment ids; the guide cites them as RC#id.")
    where = {}
    for a in reg["anchored"]:
        if a["where"].startswith("Appendix B"):
            loc = f'{a["row"]} [{a["column"]}]'
        elif a["where"] == "paragraph":
            loc = f'Heading or paragraph “{a["row"]}”'
        else:
            loc = f'Report section table: {a["row"][:110]}… [{a["column"]}]'
        where.setdefault(a["comment_id"], []).append(loc)
    t = ds.table(doc, [2.0, 3.2, 1.9, 6.6, W - 13.7], header=["RC#", "Author", "Date", "Attached to", "Text"], size=7.6)

    def cid_key(c):
        return (c["date"], int(c["id"]) if c["id"].isdigit() else 0)
    for c in sorted(reg["comments"].values(), key=cid_key):
        loc = "; ".join(where.get(c["id"], ["(reply thread)" if c["reply_to"] else "(not located)"]))
        if c["reply_to"]:
            loc = f'Reply to RC#{c["reply_to"]}. ' + loc
        ds.add_row(t, [c["id"], c["author"], c["date"], loc, c["text"]], cant_split=False)

    # ---------------- I
    h1(doc, "I  Evidence for each rule")
    note(doc, "Each rule of the guide with its status, whether it is built into the app, its sources and every quotation "
              "used as evidence. validate.py checks each quotation word for word against its source.")
    t = ds.table(doc, [2.0, 5.2, 2.4, 2.4, W - 12.0], header=["Rule", "Title", "Status", "In the app", "Sources and quotations"], size=7.6)
    for p in pol.POLICIES:
        def rid(c, p=p):
            q = c.paragraphs[0]
            ds.style_run(q.add_run(p["id"]), size=8, color=ds.NAVY, bold=True)
            ds.bookmark(q, ds.bm_name(p["id"]))

        def ev(c, p=p):
            q = c.paragraphs[0]
            ds.style_run(q.add_run("; ".join(p["sources"])), size=7.4, color=ds.MUTED)
            for k, s in p["quotes"]:
                q = c.add_paragraph(); q.paragraph_format.space_after = Pt(0.5)
                ds.style_run(q.add_run(k + ": "), size=7.4, color=ds.BLUE, bold=True)
                ds.style_run(q.add_run("“" + s + "”"), size=7.4)
            if p["impl_evidence"]:
                q = c.add_paragraph(); q.paragraph_format.space_after = Pt(0)
                ds.style_run(q.add_run("Where in the app: " + p["impl_evidence"]), size=7.2, color=ds.MUTED)
        ds.add_row(t, [rid, pol.TITLES[p["id"]], lambda c, p=p: ds.chip(c.paragraphs[0], p["status"], size=6.6),
                       ds.IMPL_LABEL[p["impl"]], ev], cant_split=False)

    # ---------------- J
    h1(doc, "J  Schema text and 4.1 calibration text")
    note(doc, "The schema text is the frozen definition shown to evaluators (app/data/definitions_v4.json, identical to "
              "SLiM_Brain/schema). The calibration text appears under “Calibration” in instrument 4.1 (app/spec.py). Both "
              "are quoted word for word.")
    t = ds.table(doc, [3.6, 3.0, 7.0, 7.0, W - 20.6], header=["Field", "Value", "Schema", "4.1 calibration text", "How they relate"], size=7.8)
    for b in pol.BASELINE:
        ds.add_row(t, [b["field"], b["value"] or "—", "“" + b["schema"] + "”",
                       ("“" + b["calibration"] + "”") if b["calibration"] else "(none)",
                       lambda c, b=b: text(c.paragraphs[0], b["reading"] + " (" + ", ".join(b["policies"]) + ")", size=7.8)])

    doc.core_properties.title = TITLE
    doc.core_properties.author = "SLiM Brain evaluation team"
    doc.save(out)


if __name__ == "__main__":
    build(json.loads((pathlib.Path(sys.argv[1]) / "register.json").read_text(encoding="utf-8")), sys.argv[2])
    print("written", sys.argv[2])
