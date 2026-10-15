"""Render the calibration guide (main document).

    python app/tools/calibration_guide/build_guide.py BUILD_DIR OUT.docx

Reads BUILD_DIR/register.json (extract_register.py) and the content model in
policies.py. Run validate.py first: this script trusts the content model.
"""
from __future__ import annotations

import collections
import json
import pathlib
import re
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import docstyle as ds  # noqa: E402
import policies as pol  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

TITLE = "SLiM Brain \u2014 Calibration Policy and Open Decisions"
SUBTITLE = "Human evaluation \u00b7 Training phase \u00b7 14 October 2026"
W = 17.8                                       # A4 portrait text width at 1.6 cm margins (cm)
STATUS_FG = ds.STATUS_STYLE["OPEN"][0]
IDS = re.compile(r"\b((?:GEN|CL|GR|FD|CO|RE|RC|CC)-\d+|D\d{1,2})\b")
POL = pol.by_id()
DEC = {d["id"]: d for d in pol.OPEN_DECISIONS}
CHK = {c["id"]: c for c in pol.CONSISTENCY_CHECKS}


def anchor_for(i: str) -> str | None:
    if i in POL:
        return ds.bm_name(i)
    if i in DEC:
        return "dec_" + i
    if i in CHK:
        return "chk_" + i.replace("-", "_")
    return None


def text(p, s, size=9, color=None, bold=None, italic=None):
    """Rich text with **bold**, _italic_ and live links on policy, decision and check ids."""
    for part in IDS.split(s):
        if not part:
            continue
        a = anchor_for(part) if IDS.fullmatch(part) else None
        if a:
            ds.link(p, part, a, size=size, color=ds.BLUE, bold=bold)
        else:
            ds.rich(p, part, size=size, color=color, bold=bold, italic=italic)
    return p


def line(cell_or_doc, s, size=9, color=None, bold=None, italic=None, after=1.5, first=False, keep=False):
    p = cell_or_doc.paragraphs[0] if first else cell_or_doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    text(p, s, size=size, color=color, bold=bold, italic=italic)
    return p


def labelled(cell, label, s, size=8.6, color=None):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    ds.style_run(p.add_run(label + " "), size=size - 0.6, color=ds.BLUE, bold=True)
    text(p, s, size=size, color=color)
    return p


def h1(doc, s, page_break=False):
    if page_break:
        ds.page_break(doc)
    return doc.add_paragraph(s, style="Heading 1")


def h2(doc, s):
    return doc.add_paragraph(s, style="Heading 2")


def status_line(p, status, impl, size=7.5):
    ds.chip(p, status, size=size)
    p.add_run(" ")
    ds.impl_tag(p, impl, size=size)


# ------------------------------------------------------------------ register helpers
class Reg:
    def __init__(self, data):
        self.rows = data["rows"]

    def get(self, claim, qid, target=None):
        for r in self.rows:
            if r["id"] == claim and r["qid"] == qid and (target is None or target in r["assigned"]):
                return r
        return None


def short_final(r):
    if r is None:
        return "unanimous, not listed in the report"
    f = r["final"] or "(no Final)"
    return f if len(f) <= 120 else f[:117] + "\u2026"


# ------------------------------------------------------------------ building blocks
def decision_of(pid):
    return [d["id"] for d in pol.OPEN_DECISIONS if pid in d["policies"]]


def policy_cards(doc, ids):
    """Cards for AGREED and PROVISIONAL rules; OPEN rules go to a compact table that links to Section 6."""
    open_ids = [i for i in ids if POL[i]["status"] == pol.OPEN]
    ids = [i for i in ids if POL[i]["status"] != pol.OPEN]
    t = ds.table(doc, [2.5, W - 2.5], size=8.3, outer=True)
    for i in ids:
        p = POL[i]

        def left(cell, p=p):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run(p["id"]), size=9.5, color=ds.NAVY, bold=True)
            ds.bookmark(q, ds.bm_name(p["id"]))
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0.5)
            ds.chip(q2, p["status"], size=6.8)
            q3 = cell.add_paragraph(); q3.paragraph_format.space_after = Pt(0)
            ds.impl_tag(q3, p["impl"], size=6.8)

        def right(cell, p=p):
            q = cell.paragraphs[0]
            q.paragraph_format.space_after = Pt(1)
            text(q, p["rule"], size=8.7)
            labelled(cell, "Apply", p["apply"], size=8.3)
            if p["exception"]:
                lab = {"AGREED": "Note", "PROVISIONAL": "Caution", "OPEN": "Open point"}[p["status"]]
                labelled(cell, lab, p["exception"], size=8.3)
            if p["status"] == pol.OPEN and decision_of(p["id"]):
                labelled(cell, "Decision", "see " + ", ".join(decision_of(p["id"])) + " in Section 6", size=8.3)
            q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(0)
            if p["examples"]:
                ds.style_run(q.add_run("Cases "), size=7.2, color=ds.BLUE, bold=True)
                ds.style_run(q.add_run(", ".join(p["examples"]) + "   "), size=7.2, color=ds.MUTED)
            ds.style_run(q.add_run("Sources "), size=7.2, color=ds.BLUE, bold=True)
            ds.style_run(q.add_run("; ".join(p["sources"])), size=7.2, color=ds.MUTED)

        ds.add_row(t, [left, right])
    if open_ids:
        open_table(doc, open_ids)
    return t


def open_table(doc, ids):
    q = ds.para(doc, "", before=4, after=2, keep=True)
    ds.style_run(q.add_run("Open in this area"), size=8.8, color=STATUS_FG, bold=True)
    t = ds.table(doc, [2.5, 7.6, W - 12.6, 2.5], header=["", "Undecided", "Meanwhile", "Decision"], size=8)
    for i in ids:
        p = POL[i]

        def left(cell, p=p):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run(p["id"]), size=9, color=ds.NAVY, bold=True)
            ds.bookmark(q, ds.bm_name(p["id"]))
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0)
            ds.chip(q2, p["status"], size=6.6)

        def mid(cell, p=p):
            q = cell.paragraphs[0]
            text(q, p["rule"].removeprefix("Undecided: "), size=8)
            q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(0)
            ds.style_run(q.add_run("Sources "), size=6.9, color=ds.BLUE, bold=True)
            ds.style_run(q.add_run("; ".join(p["sources"])), size=6.9, color=ds.MUTED)

        def meanwhile(cell, p=p):
            s = p["apply"].removeprefix("Until decided: ")
            text(cell.paragraphs[0], s[:1].upper() + s[1:], size=8)

        ds.add_row(t, [left, mid, meanwhile,
                       lambda c, p=p: text(c.paragraphs[0], ", ".join(decision_of(p["id"])) or "—", size=8.4, bold=True)])


LEGEND = {
    "AGREED": "Explicitly accepted by at least two participants at a meeting, with no unresolved dissent, or recorded as a "
              "decision in Thiago's calibration notes of 8 October (E08T), which nobody contested. Silence, a Final, or the "
              "UI alone never counts.",
    "PROVISIONAL": "Used to justify proposed Finals or written into the 4.1 calibration text, without evidence that the team "
                   "adopted it. Apply it, but expect it to be revisited.",
    "OPEN": "Deferred, listed as undecided, marked \u201c??\u201d, \u201cNo?\u201d, \u201cTo be discussed\u201d or yellow, "
            "or contradicted by another source. The guide states the competing readings; it does not choose.",
}


def legend(doc):
    t = ds.table(doc, [2.8, W - 2.8], size=8.4)
    for st in (pol.AGREED, pol.PROVISIONAL, pol.OPEN):
        def left(cell, st=st):
            ds.chip(cell.paragraphs[0], st, size=7.5)
        ds.add_row(t, [left, lambda c, st=st: text(c.paragraphs[0], LEGEND[st], size=8.4)])

    def impl_left(cell):
        q = cell.paragraphs[0]
        ds.impl_tag(q, "Implemented", size=7.5)

    ds.add_row(t, [impl_left, "Implementation is a separate label. **In UI \u2713**: verified in the code of instrument 4.1 "
                              "(calibration text or interface rule). **Not in UI**: verified absent. **\u2014**: a principle. "
                              "A rule can be implemented without being agreed, and agreed without being implemented."])
    ds.keep_table_together(t)


def dashboard(doc):
    cols = [3.4, (W - 3.4) / 3, (W - 3.4) / 3, (W - 3.4) / 3]
    t = ds.table(doc, cols, header=["Area", "AGREED", "PROVISIONAL", "OPEN"], size=8.2)
    hdr = t.rows[0]
    for k, st in enumerate((pol.AGREED, pol.PROVISIONAL, pol.OPEN), start=1):
        c = hdr.cells[k]
        c.paragraphs[0].text = ""
        ds.chip(c.paragraphs[0], st, size=7.5)
    for code, name in pol.AREAS:
        cells = [name]
        for st in (pol.AGREED, pol.PROVISIONAL, pol.OPEN):
            ps = [p for p in pol.POLICIES if p["area"] == code and p["status"] == st]

            def fill(cell, ps=ps):
                first = True
                for p in ps:
                    q = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    q.paragraph_format.space_after = Pt(1)
                    ds.link(q, p["id"], ds.bm_name(p["id"]), size=8, bold=True)
                    ds.style_run(q.add_run(" " + pol.TITLES[p["id"]]), size=8)
                    if p["impl"] == pol.IMPLEMENTED:
                        r = q.add_run(" \u00b7 UI")
                        ds.style_run(r, size=6.8, color=ds.MUTED, bold=True)
                if first:
                    ds.style_run(cell.paragraphs[0].add_run("\u2014"), size=8, color=ds.MUTED)
            cells.append(fill)
        ds.add_row(t, cells, bold=None)
    for row in t.rows[1:]:
        for p in row.cells[0].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = ds.RGBColor.from_string(ds.NAVY)


def flow_box(cell, title, body=None, fill="EAF0F8", border=ds.BLUE, italic=False):
    ds.shade(cell, fill)
    for side in ("top", "bottom", "left", "right"):
        ds.cell_border(cell, side, border, sz=6)
    q = cell.paragraphs[0]
    q.paragraph_format.space_after = Pt(0.5)
    ds.style_run(q.add_run(title), size=8.4, color=ds.NAVY, bold=True, italic=italic)
    if body:
        q2 = cell.add_paragraph()
        q2.paragraph_format.space_after = Pt(0)
        text(q2, body, size=7.8, color=ds.INK, italic=italic)


def arrow(cell, s):
    q = cell.paragraphs[0]
    q.alignment = 1
    ds.style_run(q.add_run(s), size=8, color=ds.MUTED, bold=True)


def workflow(doc, title, rows):
    """rows: list of lists of (kind, title, body) where kind in box/gate/na/arrow/blank."""
    para = ds.para(doc, "", after=2, keep=True)
    ds.style_run(para.add_run(title), size=9, color=ds.BLUE, bold=True)
    widths = [4.1, 1.3, 5.2, 1.3, 5.5]
    t = ds.table(doc, widths, size=8, rules=False, outer=False)
    for r in rows:
        cells = []
        for kind, a, b in r:
            if kind == "gate":
                cells.append(lambda c, a=a, b=b: flow_box(c, a, b, fill="DCE3EE", border=ds.NAVY))
            elif kind == "box":
                cells.append(lambda c, a=a, b=b: flow_box(c, a, b))
            elif kind == "na":
                cells.append(lambda c, a=a, b=b: flow_box(c, a, b, fill="F2F2F2", border="9AA1AA", italic=True))
            elif kind == "arrow":
                cells.append(lambda c, a=a: arrow(c, a))
            else:
                cells.append("")
        ds.add_row(t, cells)
    ds.keep_table_together(t)
    return t


# ------------------------------------------------------------------ the document
def build(reg_data, out):
    reg = Reg(reg_data)
    doc = ds.new_document(margins_cm=1.6)
    ds.header_footer(doc, "SLiM Brain \u00b7 Calibration Policy and Open Decisions", "Instrument 4.1 \u00b7 Training phase",
                     "Human evaluation \u00b7 14 October 2026")

    # ---------------- title and dashboard
    p = doc.add_paragraph(TITLE, style="Title")
    p.paragraph_format.space_after = Pt(2)
    ds.para(doc, SUBTITLE, size=11.5, color=ds.BLUE, italic=True, after=3)
    ds.para(doc, "Instrument 4.1 \u00b7 Training papers SRC-0006 and SRC-0009 \u00b7 seven evaluators \u00b7 sources up to "
                 "14 October 2026 \u00b7 companion: SLIM_BRAIN_CALIBRATION_ANNEX.docx", size=8.2, color=ds.MUTED, after=6)
    ds.para_border(doc.paragraphs[-1], "bottom", color=ds.RULE, sz=6, space=4)

    counts = collections.Counter(p["status"] for p in pol.POLICIES)
    rs = collections.Counter(r["status"] for r in reg.rows)
    intro = ds.para(doc, "", after=4)
    text(intro, "This guide states what the team has **agreed**, what has only been **applied provisionally** in proposed "
                "Final judgments, and what is still **open**, and points to the cases behind each rule. It is authoritative "
                "only to the extent of its sources: each rule names them (key in Section 8), and the annex traces every "
                "contested item of the Training report.", size=9.5)
    ds.para(doc, "1  Calibration at a glance", style="Heading 1")
    facts = ds.para(doc, "", after=4)
    text(facts, f"**{len(pol.POLICIES)} rules**: {counts['AGREED']} agreed, {counts['PROVISIONAL']} provisional, "
                f"{counts['OPEN']} open. **{len(reg.rows)} contested items** in the Training report: {rs['AGREED']} case outcomes "
                f"recorded at a meeting, {rs['PROVISIONAL']} proposed Finals, {rs['OPEN']} open. "
                f"**{len(pol.OPEN_DECISIONS)} decisions** are on the agenda (Section 6).", size=9)
    dashboard(doc)
    ds.para(doc, "Rule ids link to their entry. \u201cUI\u201d marks a rule verified in the code of instrument 4.1; it says "
                 "nothing about agreement.", size=7.6, color=ds.MUTED, before=2, after=4)
    h2(doc, "How to read the status labels")
    legend(doc)

    # ---------------- principles and workflow
    h1(doc, "2  Principles and evaluation workflow")
    principles = [
        ("Faithful mapping.", "The evaluation asks whether the Brain represents each paper faithfully for a systematic "
                              "literature mapping, not whether it finds predefined categories (GEN-1)."),
        ("Context, not literalism.", "Read the Claim as the author probably meant it, in the context of the paper, without "
                                    "going against the text (GEN-2). Textual grounding is the exception: it judges the anchors "
                                    "(GR-1 to GR-6)."),
        ("Acceptability first, attributes second.", "Relation type gates Grounding and Direction (GEN-3, agreed). Q02 gates "
                                                    "the Claim's other questions (GEN-4, provisional)."),
        ("Permissive where agreed.", "Supporting theses count as Claims (CL-1) and a defensible label is an acceptable "
                                     "Concept (CO-1). Whether field values need only be acceptable is open (GEN-11)."),
        ("A Final is a proposal.", "Unless a meeting recorded the outcome, a Final in the Training report is one reviewer's "
                                   "proposal. Agreement with a Final is not correctness, and a majority is not a "
                                   "calibration decision."),
    ]
    t = ds.table(doc, [4.3, W - 4.3], size=9)
    for a, b in principles:
        ds.add_row(t, [lambda c, a=a: ds.style_run(c.paragraphs[0].add_run(a), size=9, color=ds.NAVY, bold=True),
                       lambda c, b=b: text(c.paragraphs[0], b, size=9)])
    ds.keep_table_together(t)

    h2(doc, "Workflow and gates")
    workflow(doc, "Claim", [
        [("gate", "Q02 Central thesis or hypothesis", "Is the statement a Claim the paper advances? (CL-1)"),
         ("arrow", "Yes \u2192", None),
         ("box", "Q03\u2013Q05 Claim evaluation", "Strength and modality \u00b7 Understood on its own \u00b7 Textual grounding"),
         ("arrow", "\u2192", None),
         ("box", "Q06\u2013Q11 Schema fields", "Object \u00b7 Type \u00b7 Basis \u00b7 Jurisdiction \u00b7 Legal reference \u00b7 Temporal reference")],
        [("blank", "", None), ("arrow", "No \u2193", None),
         ("na", "Q03\u2013Q11 Not applicable", "Not asked, not stored, not counted (GEN-4, GEN-5)"),
         ("blank", "", None), ("blank", "", None)],
        [("box", "Always asked", "Q12 Concept assignment \u00b7 Q13 Candidate Concept necessary \u00b7 Q14 Missing Concepts \u00b7 the Claim's Relations (GEN-8 open)"),
         ("blank", "", None), ("blank", "", None), ("blank", "", None), ("blank", "", None)],
    ])
    merge = None
    ds.para(doc, "", after=4)
    workflow(doc, "Relation (SUPPORTS, ATTACKS, SAME_AS)", [
        [("gate", "Relation type", "Correctly classified? No if the relation should not exist or is of another kind (RE-1)"),
         ("arrow", "Yes \u2192", None),
         ("box", "Grounding", "Extracted only if one source cites the other on this point (RE-5)"),
         ("arrow", "\u2192", None),
         ("box", "Direction", "SUPPORTS and ATTACKS only; from the supporting or attacking Claim (RE-6)")],
        [("blank", "", None), ("arrow", "No \u2193", None),
         ("na", "Grounding and Direction: Not applicable", "Explain in the comment why the type is wrong (RE-1, GEN-3)"),
         ("blank", "", None), ("blank", "", None)],
    ])
    # merge the "Always asked" box across the row
    claim_flow = doc.tables[-2]
    last = claim_flow.rows[2]
    merged = last.cells[0].merge(last.cells[4])
    for extra in merged.paragraphs[2:]:
        if not extra.text.strip():
            extra._p.getparent().remove(extra._p)

    h2(doc, "Not applicable is not an answer")
    for s in ["A question skipped by a gate is **Not applicable**: it is never a Yes or a No, never a missing answer, "
              "and it is excluded from agreement statistics (GEN-5).",
              "Training answers were given before the gates existed. Stored answers are kept and read under the gates; "
              "whether to revise them is open (GEN-9, D8).",
              "In the annex, gated answers are shown but marked \u2020 and never compared with a Final."]:
        q = doc.add_paragraph(style="List Bullet")
        q.paragraph_format.space_after = Pt(1.5)
        text(q, s, size=9)

    h2(doc, "General rules")
    policy_cards(doc, ["GEN-1", "GEN-2", "GEN-3", "GEN-4", "GEN-5", "GEN-6", "GEN-10", "GEN-7", "GEN-8", "GEN-9", "GEN-11"])

    # ---------------- claims
    h1(doc, "3  Claims: eligibility, fidelity and grounding")
    policy_cards(doc, ["CL-1", "CL-2", "CL-3", "CL-4"])
    h2(doc, "Textual grounding (Q05): three different questions")
    t = ds.table(doc, [5.6, 5.6, W - 11.2], header=["Question", "Where it is judged", "Rule"], size=8.4)
    for a, b, c in [
        ("Does the paper support the Claim?", "Q02, Q03 and Q06\u2013Q11, read in the context of the paper", "GEN-2, CL-2"),
        ("Do the anchors sufficiently ground the Claim?", "Q05", "GR-1, GR-2 (sufficiency); threshold open (GR-5)"),
        ("Do the anchors reproduce every detail of the Claim?", "Not required", "GR-2: sufficiency, not completeness"),
    ]:
        ds.add_row(t, [lambda c_, a=a: ds.style_run(c_.paragraphs[0].add_run(a), size=8.4, bold=True),
                       b, lambda c_, c=c: text(c_.paragraphs[0], c, size=8.4)])
    ds.keep_table_together(t)
    ds.para(doc, "", after=2)
    policy_cards(doc, ["GR-1", "GR-2", "GR-3", "GR-4", "GR-5", "GR-6"])

    h2(doc, "Grounding decision flow")
    ds.para(doc, "The steps are a reading aid. They do not establish a precedence between rules: whether GR-4 (quotations) "
                 "or the proposed threshold of GR-5 prevails is undecided (D1).", size=8.4, color=ds.MUTED, after=3)
    steps = [
        ("1", "Locate each anchor on its page and read the surrounding passage.", "Established", "GR-1"),
        ("2", "Is the Claim's content missing from an anchor only because the 50-word limit cut the sentence?",
         "Established: not by itself a reason for In part", "GR-1"),
        ("3", "Is only a bibliographic reference missing?", "Provisional: no penalty. Reference missing mid-anchor: open",
         "GR-3"),
        ("4", "Does the Claim quote an expression that is in no anchor?", "Provisional: In part (outcome agreed for CLM-0006-015)",
         "GR-4"),
        ("5", "Is a substantive element of the Claim (a sentence, a concept) absent from the anchors?",
         "Evaluator judgement: where it is found (passage, page, paper) and whether that suffices is open", "GR-5, GR-6"),
        ("6", "Do the anchors show where the paper advances the Claim and its basis?",
         "Provisional: Yes if sufficient, even if incomplete", "GR-2"),
    ]
    t = ds.table(doc, [0.7, 8.0, 6.1, W - 14.8], header=["", "Check", "Current standing", "Rule"], size=8.3)
    for n, a, b, c in steps:
        fill = "FFFFFF"
        if b.startswith("Evaluator"):
            fill = "FBF3F4"

        def st(cell, b=b):
            q = cell.paragraphs[0]
            key = b.split(":")[0]
            col = {"Established": "1D5B36", "Provisional": "6E4300", "Evaluator judgement": "8A1C2B"}.get(key, ds.INK)
            ds.style_run(q.add_run(key), size=8.3, color=col, bold=True)
            rest = b[len(key):]
            if rest:
                text(q, rest, size=8.3)
        ds.add_row(t, [lambda c_, n=n: ds.style_run(c_.paragraphs[0].add_run(n), size=9, color=ds.NAVY, bold=True),
                       a, st, lambda c_, c=c: text(c_.paragraphs[0], c, size=8.3)], fill=fill)
    ds.keep_table_together(t)

    h2(doc, "Grounding cases")
    t = ds.table(doc, [2.3, 2.6, 7.0, 2.4, W - 14.3], header=["Claim", "Answers", "Final (verbatim)", "Status", "Rule"], size=8.1)
    for claim, rule in [("CLM-0006-004", "GR-4"), ("CLM-0006-015", "GR-4, GR-5"), ("CLM-0006-018", "GR-2"),
                        ("CLM-0009-006", "GR-1, GR-6"), ("CLM-0009-010", "GR-6"), ("CLM-0009-011", "GR-1, GR-6")]:
        r = reg.get(claim, "Q05")
        ds.add_row(t, [claim, r["split_stored"], short_final(r),
                       lambda c_, r=r: ds.chip(c_.paragraphs[0], r["status"], size=7),
                       lambda c_, rule=rule: text(c_.paragraphs[0], rule, size=8.1)])

    # ---------------- fields
    h1(doc, "4  Schema fields (Q06\u2013Q11)")
    box = ds.table(doc, [W], size=8.8, rules=False)

    def acc(cell):
        ds.shade(cell, ds.PANEL)
        q = cell.paragraphs[0]
        ds.style_run(q.add_run("Acceptable or uniquely preferable?  "), size=9, color=ds.NAVY, bold=True)
        text(q, "Each field question asks whether the value is \u201ccorrectly assigned according to the schema definition\u201d. "
                "The team has not decided whether Yes means that the value is acceptable or that it is the best one "
                "(GEN-11, D2). Until it does, answer No only when the assigned value is not defensible, and name the value "
                "you would choose in the comment. Mixed Claims: judge the dominant value (GEN-6).", size=8.8)
    ds.add_row(box, [acc])
    ds.para(doc, "", after=2)
    policy_cards(doc, ["FD-1", "FD-2", "FD-3", "FD-4", "FD-5", "FD-6", "FD-7", "FD-8", "FD-9", "FD-10", "FD-11"])
    note = ds.para(doc, "", before=3, after=2)
    text(note, "Where the 4.1 calibration text departs from the frozen schema (Claim object law and other; Basis literature), "
               "this guide treats it as an interpretation, not a schema change. Annex J sets the schema text and the "
               "calibration text side by side.", size=8.3, color=ds.MUTED)

    # ---------------- concepts and relations
    h1(doc, "5  Concepts and relations")
    h2(doc, "Concepts: three separate questions")
    t = ds.table(doc, [3.6, 6.2, 4.6, W - 14.4], header=["Question", "Ask", "Not a reason for No", "Rules"], size=8.3)
    for a, b, c, d in [
        ("Q12 Does the assigned Concept fit?", "Is it a defensible label for what the Claim is about, even partly, and does it capture a substantive aspect?",
         "The Concept's family; the Claim not naming it", "CO-1, CO-2, CO-3"),
        ("Q13 Is the candidate necessary?", "Without it, could the existing vocabulary trace what the Claim is about?",
         "That it fits (fit is Q12)", "CO-4"),
        ("Q14 Is a Concept missing?", "Is there a clear omission whose absence is felt, especially from the existing vocabulary?",
         "That another relevant Concept could be added", "CO-5"),
    ]:
        ds.add_row(t, [lambda c_, a=a: ds.style_run(c_.paragraphs[0].add_run(a), size=8.3, bold=True), b, c,
                       lambda c_, d=d: text(c_.paragraphs[0], d, size=8.3)])
    ds.keep_table_together(t)
    ds.para(doc, "In the SRC-0009 Finals, many proposed Missing Concepts are called \u201creasonable\u201d. The annex records "
                 "these as considered reasonable, not as adopted or necessary (CC-4).", size=8.4, color=ds.MUTED, before=2, after=3)
    policy_cards(doc, ["CO-1", "CO-2", "CO-3", "CO-4", "CO-5", "CO-6", "CO-7", "CO-8"])

    h2(doc, "Relations: what to record")
    t = ds.table(doc, [5.4, 2.0, 2.2, 2.4, W - 12.0], header=["Situation", "Relation type", "Grounding", "Direction", "Comment"], size=8.2)
    for row in [
        ("No meaningful relationship: the edge should not exist", "No", "N/A", "N/A", "Say that the edge should not exist (RE-1)"),
        ("A relationship exists, but not of the extracted type (e.g. compatible findings typed SUPPORTS)", "No", "N/A", "N/A", "Name the relationship you see (RE-1)"),
        ("Correct type, wrong direction", "Yes", "judge", "No", "RE-2"),
        ("Correct type, wrong grounding", "Yes", "No", "judge", "RE-5"),
        ("SAME_AS, correct", "Yes", "judge", "not asked", "RE-6"),
    ]:
        ds.add_row(t, [row[0], lambda c_, v=row[1]: ds.style_run(c_.paragraphs[0].add_run(v), size=8.2, bold=True),
                       row[2], row[3], lambda c_, v=row[4]: text(c_.paragraphs[0], v, size=8.2)])
    ds.keep_table_together(t)
    ds.para(doc, "The first two situations produce the same answers; only the comment distinguishes them. Error analysis "
                 "of relations therefore depends on comments being written.", size=8.4, color=ds.MUTED, before=2, after=3)
    policy_cards(doc, ["RE-1", "RE-2", "RE-3", "RE-4", "RE-5", "RE-6"])

    h2(doc, "Disputed relations")
    t = ds.table(doc, [4.6, 4.2, 4.2, W - 13.0], header=["Relation", "Relation type (Final)", "Grounding / Direction (Final)", "Standing"], size=8)
    for src, tgt, standing in [
        ("CLM-0009-004", "CLM-0011-006", "Open: tentative No; parallel findings (D5)"),
        ("CLM-0009-004", "CLM-0016-004", "Open: Final text truncated (D5)"),
        ("CLM-0009-004", "CLM-0018-011", "Open: \u201c??\u201d (D5)"),
        ("CLM-0009-006", "CLM-0008-009", "Type: provisional No. Direction: Final N/A contradicted by its own comment (CC-6)"),
        ("CLM-0009-012", "CLM-0016-007", "Provisional No without stated reason; 4 of 7 answered Yes (CC-10)"),
        ("CLM-0009-013", "CLM-0019-009", "Open: \u201c??\u201d on all three questions (D5)"),
    ]:
        rt = reg.get(src, "REL_TYPE", tgt)
        gr = reg.get(src, "REL_GROUNDING", tgt)
        di = reg.get(src, "REL_DIRECTION", tgt)
        typ = (rt["assigned"].split(" ")[0] if rt else (di or gr)["assigned"].split(" ")[0])

        if rt and rt["final_value"] == "No" and rt["status"] != "OPEN":
            absent = "not applicable (type No)"
        elif rt and rt["status"] == "OPEN":
            absent = "unanimous in stored answers; applies only if type is Yes"
        else:
            absent = "unanimous, not listed in the report"

        def gd(cell, gr=gr, di=di, absent=absent):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run("G: "), size=8, color=ds.MUTED, bold=True)
            ds.style_run(q.add_run(short_final(gr) if gr else absent), size=8)
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0)
            ds.style_run(q2.add_run("D: "), size=8, color=ds.MUTED, bold=True)
            ds.style_run(q2.add_run(short_final(di) if di else absent), size=8)
        ds.add_row(t, [f"{src} \u2192 {tgt}\n{typ} (inferred)", f'{short_final(rt)}' + (f'\n[{rt["split_stored"]}]' if rt else ""),
                       gd, lambda c_, s=standing: text(c_.paragraphs[0], s, size=8)])

    # ---------------- open decisions
    h1(doc, "6  Open decisions requiring resolution")
    ds.para(doc, "Ordered by priority. Priority 1 affects how comparable the evaluators' answers are; priority 2 affects "
                 "what enters the statistics; priority 3 affects reporting only. Ownership is not assigned in any source.",
            size=8.6, color=ds.MUTED, after=4)
    t = ds.table(doc, [1.9, 8.4, W - 10.3], header=["", "Decision and competing readings", "Cases \u00b7 why it matters \u00b7 to decide"], size=8.1)
    for d in sorted(pol.OPEN_DECISIONS, key=lambda x: (x["priority"], int(x["id"][1:]))):
        def head(cell, d=d):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run(d["id"]), size=10.5, color=ds.NAVY, bold=True)
            ds.bookmark(q, "dec_" + d["id"])
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0.5)
            ds.style_run(q2.add_run(f'Priority {d["priority"]}'), size=7, color=ds.MUTED, bold=True)
            q3 = cell.add_paragraph(); q3.paragraph_format.space_after = Pt(0)
            text(q3, ", ".join(d["policies"]), size=7.2)

        def body(cell, d=d):
            q = cell.paragraphs[0]
            q.paragraph_format.space_after = Pt(1)
            ds.style_run(q.add_run(d["title"]), size=8.9, color=ds.NAVY, bold=True)
            q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(1.5)
            text(q, d["question"], size=8.1)
            for s_ in d["interpretations"]:
                b = cell.add_paragraph(); b.paragraph_format.space_after = Pt(0.5)
                b.paragraph_format.left_indent = Cm(0.3); b.paragraph_format.first_line_indent = Cm(-0.3)
                text(b, "\u2013 " + s_, size=7.9)

        def side(cell, d=d):
            q = cell.paragraphs[0]
            q.paragraph_format.space_after = Pt(1)
            if d["examples"]:
                ds.style_run(q.add_run("Cases "), size=7.4, color=ds.BLUE, bold=True)
                ds.style_run(q.add_run(", ".join(d["examples"])), size=7.4, color=ds.MUTED)
                q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(1)
            ds.style_run(q.add_run("Why "), size=7.8, color=ds.BLUE, bold=True)
            text(q, d["why"], size=7.9)
            q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(0)
            ds.style_run(q.add_run("To decide "), size=7.8, color=ds.BLUE, bold=True)
            text(q, d["decide"], size=7.9)
        ds.add_row(t, [head, body, side])

    # ---------------- implications
    h1(doc, "7  Methodological implications and consistency checks")
    h2(doc, "Claim recall: four different things")
    ds.para(doc, "Translating a paper into a finite set of Claims is indeterminate (Vaclav, E09V, citing Quine); two experts "
                 "need not produce the same set. The proposal under discussion (RC-1) therefore does not measure absolute "
                 "recall.", size=8.8, after=3)
    t = ds.table(doc, [4.0, 5.6, W - 9.6], header=["Measure", "What it needs", "In this evaluation"], size=8.3)
    for a, b, c in [
        ("Conventional recall", "A defined reference set of Claims per paper", "Not available, and not sought (RC-1)"),
        ("Holistic coverage", "Each evaluator's judgement of how completely the set represents the paper",
         "The recall scale (None / Some / Most / All); answered by everyone"),
        ("Expert-proposed additions", "Missing Claims written as propositions with a page",
         "Collected when recall is below All; no Final, by design (E13A)"),
        ("Validation of additions", "Others judge whether each addition is a meaningful improvement to the set",
         "Proposed (Vaclav, Alessandro); not adopted or designed (D9)"),
    ]:
        ds.add_row(t, [lambda c_, a=a: ds.style_run(c_.paragraphs[0].add_run(a), size=8.3, bold=True), b,
                       lambda c_, c=c: text(c_.paragraphs[0], c, size=8.3)])
    ds.keep_table_together(t)
    ds.para(doc, "Restatements: four evaluators grouped SRC-0006 Claims in four different ways; whether a restatement means "
                 "the same assertion or the same topic is open (RC-2).", size=8.6, before=3, after=2)
    policy_cards(doc, ["RC-1", "RC-2"])

    h2(doc, "Consequences for agreement statistics")
    for s in [
        "Gated answers are excluded, so the number of rated pairs differs between questions; report it with each coefficient (GEN-5).",
        "Finals are mostly proposals by one reviewer. Agreement with a Final measures convergence, not correctness.",
        "A majority among seven evaluators is not a calibration decision; the register never uses one to resolve an open item.",
        "Answers revised after reading the Finals are no longer independent; agreement before and after calibration should be reported separately.",
        "Thibault (E11S): the results describe the current models and may not transfer to later ones; this belongs in the Discussion.",
    ]:
        q = doc.add_paragraph(style="List Bullet")
        q.paragraph_format.space_after = Pt(1.5)
        text(q, s, size=8.8)
    ds.para(doc, "The last three points are editorial observations of this guide, not team decisions.", size=7.8,
            color=ds.MUTED, after=3)

    h2(doc, "Consistency checks")
    t = ds.table(doc, [1.4, 3.0, W - 7.4, 3.0], header=["", "Items", "Point", "Rules"], size=8)
    for c in pol.CONSISTENCY_CHECKS:
        def cid(cell, c=c):
            q = cell.paragraphs[0]
            ds.style_run(q.add_run(c["id"]), size=8.2, color=ds.NAVY, bold=True)
            ds.bookmark(q, "chk_" + c["id"].replace("-", "_"))
            q2 = cell.add_paragraph(); q2.paragraph_format.space_after = Pt(0)
            ds.style_run(q2.add_run(c["origin"]), size=6.8, color=ds.MUTED)
        ds.add_row(t, [cid, ", ".join(c["items"]) or "\u2014", lambda c_, c=c: text(c_.paragraphs[0], c["point"], size=8),
                       lambda c_, c=c: text(c_.paragraphs[0], ", ".join(c["policies"]), size=8)])

    # ---------------- process and sources
    h1(doc, "8  Using the annex; sources")
    for s in [
        "**Review your answers.** E08T asks each evaluator to revise where they \u201ccontradicted the majority\u201d; Vaclav "
        "(E09V) proposes checking the items where one's answer differs from the Final. Annex F lists these items per evaluator. "
        "Items whose Final is open are listed separately and are not corrections.",
        "**Trace a case.** Annex B gives every contested item with the seven answers, the verbatim Final, its normalised value, "
        "its status and the linked rules. Annex H reproduces all reviewer comments.",
        "**Annex I** repeats every rule with all of its quotations, so that each status can be checked against its evidence.",
    ]:
        q = doc.add_paragraph(style="List Bullet")
        q.paragraph_format.space_after = Pt(1.5)
        text(q, s, size=8.8)
    h2(doc, "Source key")
    short = [("M28", "Meeting, 28 Sep (instrument 4.0)"), ("M07", "Meeting, 7 Oct (4.1 Training answers)"),
             ("E08T", "Thiago, 8 Oct: calibration notes"), ("E08A", "Alessandro, 8 Oct: gates"),
             ("E09T", "Thiago, 9 Oct: Claim gate"), ("E09V", "Vaclav, 9 Oct: Finals SRC-0006; recall"),
             ("E11S", "Thibault, 11 Oct: model change"), ("E12T", "Thiago, 12 Oct: UI ready"),
             ("E13V", "Vaclav, 13 Oct: resolve open points"), ("E13A", "Alessandro, 13 Oct: recall"),
             ("REP", "Training report, 14 Oct (Appendix B, Finals)"), ("RC#n", "Reviewer comment n (Annex H)"),
             ("F:\u2026", "Final cell: Claim / question [/ target]"), ("EC:\u2026", "Evaluator's own comment in the report"),
             ("SPEC", "Instrument 4.1 code and calibration texts"), ("DEF", "Frozen schema shown to evaluators"),
             ("SUM", "Earlier summary (secondary only)"), ("", "")]
    t = ds.table(doc, [1.5, W / 2 - 1.5, 1.5, W / 2 - 1.5], size=7.8)
    for i in range(0, len(short), 2):
        (k1, v1), (k2, v2) = short[i], short[i + 1]
        ds.add_row(t, [lambda c_, k=k1: ds.style_run(c_.paragraphs[0].add_run(k), size=7.8, color=ds.NAVY, bold=True), v1,
                       lambda c_, k=k2: ds.style_run(c_.paragraphs[0].add_run(k), size=7.8, color=ds.NAVY, bold=True), v2])
    ds.para(doc, "Meeting references give the transcript timestamp. In M28 several speakers share the label \u201cGiuseppe "
                 "Pisano\u201d (a room microphone); attributions there are uncertain. The email thread is available as "
                 "Calibration/email-conversation.txt. Full source descriptions: Annex A.", size=7.8, color=ds.MUTED, before=3)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE
    doc.core_properties.author = "SLiM Brain evaluation team"
    doc.save(out)


if __name__ == "__main__":
    build(json.loads((pathlib.Path(sys.argv[1]) / "register.json").read_text(encoding="utf-8")), sys.argv[2])
    print("written", sys.argv[2])
