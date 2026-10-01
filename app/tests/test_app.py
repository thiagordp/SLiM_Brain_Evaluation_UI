"""Checks for the evaluation application, instrument 4.1.

    python app/tests/test_app.py            # everything
    python app/tests/test_app.py adapter    # only tests whose name contains "adapter"

Storage is an in-memory Google Sheets workbook (`tests/fakes.py`); nothing here
touches a real workbook. The UI checks run the real Streamlit script headlessly
with `streamlit.testing.v1.AppTest`.
"""
from __future__ import annotations

import contextlib
import html
import itertools
import io
import json
import os
import pathlib
import re
import shutil
import sys
import tempfile
import traceback

APP = pathlib.Path(__file__).resolve().parents[1]
sys.path[:0] = [str(APP), str(APP / "tests"), str(APP / "tools")]
os.environ["HE_APP_PASSWORD"] = "test-password"
os.environ["HE_ADMIN_SECRET"] = "test-admin"
for name in ("GOOGLE_SHEET_ID", "HE_GOOGLE_SHEET_ID", "GOOGLE_SERVICE_ACCOUNT_JSON",
             "HE_GOOGLE_CREDENTIALS", "GOOGLE_CREDENTIALS", "HE_ROUND_ID", "HE_BRAIN_DIR"):
    os.environ.pop(name, None)

import allocation  # noqa: E402
import analysis  # noqa: E402
import bootstrap_round  # noqa: E402
import brain as brain_module  # noqa: E402
import conceptsearch  # noqa: E402
import fakes  # noqa: E402
import freeze_definitions  # noqa: E402
import manifest  # noqa: E402
import preflight  # noqa: E402
import progress  # noqa: E402
import sheets  # noqa: E402
import spec  # noqa: E402
import store  # noqa: E402

ROUND = manifest.round_id()
BRAIN = brain_module.load_brain()
FAILURES: list[str] = []
CHECKS = 0


# ------------------------------------------------ choices from the allocation
#: Tests pick their evaluator and papers from the generated manifest, so a new
#: allocation never needs the tests rewritten. Thiago is used because he is in
#: every allocation so far and is the administrator.
ASSIGNMENTS = manifest.load_manifest_file()["assignments"]
EVALUATOR = "thiago"


def _assigned(evaluator: str, phase: str) -> list[str]:
    return sorted(a["source_id"] for a in ASSIGNMENTS
                  if a["evaluator_id"] == evaluator and a["phase_id"] == phase)


def _rich(source: str) -> bool:
    """A paper exercising every Claim-page part: Datasets, several Concepts on
    its first Claim, and a Relation starting from that Claim."""
    first = BRAIN.claims_of(source)[0]["id"]
    return (bool(BRAIN.dataset_ids_of(source)) and len(BRAIN.concepts_of_claim(first)) >= 2
            and bool(BRAIN.relations_from_claim(first)))


AGR_SOURCE = next(s for s in _assigned(EVALUATOR, "agreement") if _rich(s))
AGR_PARTNER = next(a["evaluator_id"] for a in ASSIGNMENTS
                   if a["phase_id"] == "agreement" and a["source_id"] == AGR_SOURCE
                   and a["evaluator_id"] != EVALUATOR)
IND_SOURCE = _assigned(EVALUATOR, "individual")[0]
NO_DATASET_SOURCE = next(s for s in _assigned(EVALUATOR, "individual")
                         + _assigned(EVALUATOR, "agreement")
                         if not BRAIN.dataset_ids_of(s))
NO_DATASET_PHASE = ("individual" if NO_DATASET_SOURCE in _assigned(EVALUATOR, "individual")
                    else "agreement")


def check(condition, message: str) -> None:
    global CHECKS
    CHECKS += 1
    if not condition:
        FAILURES.append(message)
        print(f"  FAIL  {message}")


def fresh_round() -> fakes.FakeWorkbook:
    """A fake workbook bootstrapped exactly as a real round would be."""
    book = fakes.FakeWorkbook()
    store.use_workbook(book)
    manifest.invalidate()
    sys.argv = ["bootstrap_round.py"]
    with contextlib.redirect_stdout(io.StringIO()):
        code = bootstrap_round.main()
    assert code == 0, "bootstrap failed"
    store.forget_all()
    return book


def rid(phase: str, evaluator: str, source: str) -> str:
    return store.review_id(ROUND, phase, evaluator, source)


def start(review_id: str) -> None:
    store.start_review(review_id, True, eval_spec_version=spec.EVAL_SPEC_VERSION)


def answer_everything(review_id: str, source_id: str, *, recall=spec.RECALL_ALL) -> None:
    """Complete one review directly through the store, as fast as a test can."""
    data = store.load_review(review_id)
    records = []
    for row in data.responses.values():
        question = spec.BY_KEY[row["question_key"]]
        if question is spec.Q14:
            answer = spec.Q14_NONE_MISSING
        elif question is spec.RECALL:
            answer = recall
        elif question is spec.MISSING_CLAIMS:
            answer = "" if recall == spec.RECALL_ALL else "A missing thesis (p. 3)."
        elif question is spec.RESTATEMENTS:
            answer = spec.RESTATEMENTS_NONE
        else:
            answer = spec.YES
        records.append({"tab": sheets.RESPONSES, "key": row["response_key"],
                        "values": {**row, "answer": answer}, "rid": review_id})
    for row in data.concepts.values():
        records.append({"tab": sheets.CONCEPT_RESPONSES, "key": row["response_key"],
                        "values": {**row, "answer": spec.YES}, "rid": review_id})
    for row in data.relations.values():
        directed = spec.DIRECTION.applies_to({"type": row["relation_type"]})
        records.append({"tab": sheets.RELATION_RESPONSES, "key": row["response_key"],
                        "values": {**row, "grounding_answer": spec.YES,
                                   "direction_answer": spec.YES if directed else "",
                                   "type_answer": spec.YES}, "rid": review_id})
    store.save_many(records)


def brain_copy() -> pathlib.Path:
    target = pathlib.Path(tempfile.mkdtemp(prefix="he-brain-"))
    shutil.copytree(BRAIN.root / "wiki", target / "wiki")
    shutil.copytree(BRAIN.root / "schema", target / "schema")
    return target


# =================================================================== adapter
def _on_disk():
    """Counts taken straight from the Brain files, independently of the adapter,
    so an ingest that adds Sources changes the expectation with it."""
    wiki = BRAIN.root / "wiki"
    lines = lambda p: [json.loads(l) for l in p.read_text().splitlines() if l.strip()]
    claims = lines(wiki / "claims" / "claims.jsonl")
    edges = lines(wiki / "graph" / "edges.jsonl")
    claim_ids = {c["id"] for c in claims}
    return {
        "sources": len(list((wiki / "sources").glob("SRC-*.md"))),
        "datasets": len(list((wiki / "datasets").glob("DST-*.md"))),
        "concepts": {p.stem for p in (wiki / "concepts").glob("CPT-*.md")},
        "claims": len(claims),
        "relations": sum(1 for e in edges if e["type"] in ("SUPPORTS", "ATTACKS", "SAME_AS")
                         and e["from"] in claim_ids and e["to"] in claim_ids),
    }


def test_adapter_reads_current_schema():
    disk = _on_disk()
    check(len(BRAIN.sources) == disk["sources"] and len(BRAIN.claims) == disk["claims"],
          f"adapter: all {disk['sources']} Sources and {disk['claims']} Claims")
    check(len(BRAIN.datasets) == disk["datasets"] and not BRAIN.load_problems,
          "adapter: every Dataset page is read, including those YAML would reject")
    check(":" in BRAIN.dataset("DST-0004")["description"],
          "adapter: a description containing a colon is read whole")
    check(BRAIN.dataset_ids_of("SRC-0011") == ["DST-0008", "DST-0009", "DST-0010"],
          "adapter: Datasets come from the plural Claim field")
    check(BRAIN.dataset_ids_of("SRC-0003") == [], "adapter: SRC-0003 has no Dataset")
    relations = BRAIN.claim_relations()
    hosted = sum(len(BRAIN.relations_hosted(s)) for s in BRAIN.source_ids)
    check(hosted == len(relations) == disk["relations"],
          "adapter: every Claim Relation is hosted by exactly one Source")
    check(all(r["type"] in ("SUPPORTS", "ATTACKS", "SAME_AS") for r in relations),
          "adapter: only current Claim-to-Claim relation types")
    check(brain_module.concept_label("CPT-irac-analysis") == "IRAC analysis"
          and brain_module.concept_label("CPT-access-to-justice") == "Access to justice",
          "adapter: Concept names are derived deterministically from ids")
    vocabulary = BRAIN.vocabulary(spec.concept_grid())
    grid_only = [v for v in vocabulary if not v["has_page"]]
    grid = {c for ids in spec.concept_grid().values() for c in ids}
    check(len(grid_only) == len(grid - disk["concepts"])
          and all(v["definition"] == "" for v in grid_only),
          "adapter: grid-only anchors are in the vocabulary, with no invented definition")
    check(BRAIN.canonical_schema_version == "0.1.0", "adapter: canonical schema version")
    parsed = brain_module.parse_frontmatter(
        "---\nid: X\nlist: [a, b]\njson: [\"c, d\"]\nblock:\n  - e\n  - f\nnote: a: b\n---\n")
    check(parsed == {"id": "X", "list": ["a", "b"], "json": ["c, d"],
                     "block": ["e", "f"], "note": "a: b"},
          "adapter: frontmatter is read the way the Brain's own tools read it")


# =============================================================== definitions
def test_definitions_are_frozen_verbatim():
    skills = freeze_definitions.DEFAULT_SKILLS
    if skills.exists():
        rebuilt = freeze_definitions.build(BRAIN.root / "schema", skills)
        check(rebuilt == spec.definitions(),
              "definitions: the frozen file matches the schema and skills on disk")
    for key, entry in spec.definitions()["entries"].items():
        path = (BRAIN.root / entry["source"]) if entry["source"].startswith("schema/") \
            else skills.parent.parent / entry["source"]
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        check(entry["text"] in text, f"definitions: {key} text is verbatim")
        for item in entry.get("items", []):
            if item.get("text"):
                check(item["text"] in text, f"definitions: {key} / {item.get('value')} verbatim")
            for sub in item.get("sub", []):
                check(sub in text, f"definitions: {key} sub-item verbatim")
    grounding = spec.definition("edge.grounding")
    check(grounding["source"] == ".claude/skills/create-edges/SKILL.md"
          and grounding["items"][1]["text"] == "otherwise.",
          "definitions: Grounding comes word for word from the create-edges skill")
    check(spec.plain("one of `public` | `on_request`") == "one of public | on_request",
          "definitions: display drops only the code markup")


# ====================================================================== spec
AGREED = {
    "CLAIM_Q02_CENTRAL_THESIS": "Does this statement represent a central thesis or "
                                "hypothesis that the paper advances as part of its "
                                "contribution?",
    "CLAIM_Q03_MODALITY": "Does the Claim preserve the strength and modality of the source?",
    "CLAIM_Q04_STANDALONE": "Can this Claim be understood on its own, without missing "
                            "context from the paper?",
    "CLAIM_Q05_GROUNDING": "Do the anchors provide sufficient textual grounding for this "
                           "Claim in the source?",
    "CLAIM_Q06_CLAIM_OBJECT": "Is Claim object correctly assigned according to the schema "
                              "definition?",
    "CLAIM_Q07_CLAIM_TYPE": "Is Claim type correctly assigned according to the schema "
                            "definition?",
    "CLAIM_Q08_BASIS": "Is Basis correctly assigned according to the schema definition?",
    "CLAIM_Q09_CLAIM_JURISDICTION": "Is Claim jurisdiction correctly assigned according "
                                    "to the schema definition?",
    "CLAIM_Q10_LEGAL_REFERENCE": "Is Legal reference correctly assigned according to the "
                                 "schema definition?",
    "CLAIM_Q11_TEMPORAL_REFERENCE": "Is Temporal reference correctly assigned according "
                                    "to the schema definition?",
    "CLAIM_Q12_CONCEPT": "Is this Concept correctly assigned to the Claim?",
    "CLAIM_Q13_CANDIDATE": "Is this candidate Concept necessary for representing this "
                           "Claim, given the existing Concept vocabulary?",
    "CLAIM_Q14_MISSING_CONCEPTS": "Are any mapping-relevant Concepts missing from this Claim?",
    "REL_GROUNDING": "Is the Grounding value correctly assigned according to the schema "
                     "definition?",
    "REL_DIRECTION": "Is the direction of this relation correct?",
    "REL_TYPE": "Is the Relation type correct for the relationship between these two Claims?",
    "DATASET_NODE": "Does this record represent a dataset, benchmark or corpus that a "
                    "Claim in this Source actually rests on?",
    "DATASET_INTRODUCED_BY": "Is Introduced by correctly assigned according to the schema "
                             "definition?",
    "DATASET_LANGUAGE": "Is Language correctly assigned according to the schema definition?",
    "DATASET_JURISDICTION": "Is Jurisdiction correctly assigned according to the schema "
                            "definition?",
    "DATASET_DESCRIPTION": "Does the Dataset description correctly and sufficiently "
                           "represent the Dataset characteristics reported by the Source?",
    "DATASET_AVAILABILITY": "Is Availability correctly assigned according to the schema "
                            "definition?",
    "SOURCE_CLAIM_RECALL": "How completely does the current set of extracted Claims "
                           "represent the central theses or hypotheses that this Source "
                           "advances as part of its contribution?",
    "SOURCE_MISSING_CLAIMS": "Which central Claims are missing from the extracted set?",
    "SOURCE_RESTATEMENTS": "Are any extracted Claims restatements of the same proposition?",
}


def test_spec_wording_and_semantics():
    check(set(AGREED) == set(spec.BY_KEY), "spec: exactly the agreed questions exist")
    for key, text in AGREED.items():
        check(spec.BY_KEY[key].text == text, f"spec: {key} uses the agreed wording")
    check(spec.EVAL_SPEC_VERSION == "4.1", "spec: the instrument is version 4.1")
    check("CLAIM_Q01_RESTATEMENT" not in spec.BY_KEY
          and not any(k.startswith("CLAIM_Q01") for k in spec.BY_KEY),
          "spec: the per-Claim restatement question is gone")
    check([q.key for q in spec.claim_scalar_questions(spec.PART_CLAIM)]
          == ["CLAIM_Q02_CENTRAL_THESIS", "CLAIM_Q03_MODALITY", "CLAIM_Q04_STANDALONE",
              "CLAIM_Q05_GROUNDING"], "spec: Q2–Q5 keep their keys")
    check(spec.RESTATEMENTS.options == (spec.RESTATEMENTS_NONE, spec.RESTATEMENTS_PRESENT)
          and spec.RESTATEMENT_LABEL == {"none": "No", "present": "Yes"}
          and spec.problem(spec.RESTATEMENTS, "present")
          and not spec.problem(spec.RESTATEMENTS, "none"),
          "spec: Restatements are stored as none / present and shown as No / Yes")
    check([q.key for q in spec.relation_questions()]
          == ["REL_GROUNDING", "REL_DIRECTION", "REL_TYPE"],
          "spec: Relations ask Grounding, Direction, Relation type in that order")
    check([q.key for q in spec.relation_questions({"type": "SAME_AS"})]
          == ["REL_GROUNDING", "REL_TYPE"]
          and len(spec.relation_questions({"type": "SUPPORTS"})) == 3
          and len(spec.relation_questions({"type": "ATTACKS"})) == 3,
          "spec: Direction is asked for SUPPORTS and ATTACKS, never SAME_AS")
    check(spec.DIRECTION.options == spec.BINARY
          and spec.DIRECTION.optional_comment_on == (spec.NO,),
          "spec: Direction is Yes / No with an optional comment on No")
    frozen = json.dumps(spec.definitions(), ensure_ascii=False)
    check(not any(line in frozen for lines in spec.CALIBRATION.values() for line in lines),
          "spec: calibration text is not in the frozen definitions")
    check(spec.CALIBRATION["claim.general"] == (
        "Evaluate this Claim from its statement and anchors.",
        "Judge the Claim fields from what the anchors support.",
        "Restatements are assessed in Claim recall."),
        "spec: the Claim-level calibration has exactly the three agreed lines")
    q6 = spec.BY_KEY["CLAIM_Q06_CLAIM_OBJECT"]
    check(not spec.problem(q6, spec.YES) and spec.problem(q6, spec.NO),
          "spec: Questions 6–11 Yes means correct")
    check(spec.BY_KEY["CLAIM_Q05_GROUNDING"].options == spec.TERNARY,
          "spec: Question 5 is Yes / In part / No")
    check(spec.RECALL.options == spec.RECALL_LEVELS and len(spec.RECALL_LEVELS) == 4,
          "spec: Claim recall has four ordered levels")
    check(spec.ordinal(spec.RECALL, spec.RECALL_NONE) == 0
          and spec.ordinal(spec.RECALL, spec.RECALL_ALL) == 3, "spec: recall order")
    description = spec.BY_KEY["DATASET_DESCRIPTION"]
    check(spec.comment_required(description, spec.IN_PART)
          and spec.comment_required(description, spec.NO)
          and not spec.comment_required(q6, spec.NO),
          "spec: only the Description comment is required")
    everything = json.dumps([dataclass_text(q) for q in spec.QUESTIONS])
    check(not re.search(r"HE-\d", everything), "spec: no HE identifier is exposed")
    check("`" not in everything, "spec: no code typography in question wording")


def dataclass_text(question):
    return [question.text, question.title, question.comment_help, question.comment_label]


# ================================================================ allocation
def test_allocation_rules():
    data = allocation.load(manifest.allocation_path())
    check(allocation.validate(data, BRAIN) == [], "allocation: the allocation file is valid")
    evaluators, rows = allocation.expand(data, BRAIN)
    phases = {p: [r for r in rows if r["phase_id"] == p] for p in allocation.PHASES}
    training = {r["source_id"] for r in phases["training"]}
    check(training == {"SRC-0006", "SRC-0009"}
          and len(phases["training"]) == 2 * len(evaluators),
          "allocation: the 2 agreed Training papers, for every evaluator")
    groups = data["pairs"]
    check(all(len(members) >= 2 for members in groups.values())
          and sorted(m for ms in groups.values() for m in ms)
          == sorted(e["evaluator_id"] for e in evaluators),
          "allocation: every evaluator belongs to exactly one group of two or more")
    for split, spec_ in data["agreement"].items():
        size = len(groups[spec_["pair"]])
        for entry in spec_["sources"]:
            made = [r for r in phases["agreement"] if r["source_id"] == entry["source"]]
            check(len(made) == size,
                  f"allocation: {split} {entry['source']} is evaluated by every group member")
    agreement_papers = {r["source_id"] for r in phases["agreement"]}
    remaining = len(BRAIN.sources) - len(training) - len(agreement_papers)
    check(len(phases["individual"]) == remaining,
          f"allocation: the other {remaining} papers are Individual")
    per = sorted(sum(1 for r in phases["individual"] if r["evaluator_id"] == e["evaluator_id"])
                 for e in evaluators)
    check(per[-1] - per[0] <= 1 and sum(per) == remaining,
          "allocation: Individual papers as even as possible")

    def broken(mutate):
        copy = json.loads(json.dumps(data))
        mutate(copy)
        return allocation.validate(copy, BRAIN)

    check(any("Training Source" in p for p in broken(
        lambda d: d["agreement"][next(iter(d["agreement"]))]["sources"].append("SRC-0006"))),
        "allocation: a Training Source used elsewhere is refused")
    check(any("more than once" in p for p in broken(
        lambda d: d["individual"][next(iter(d["individual"]))]["sources"].append(
            d["agreement"][next(iter(d["agreement"]))]["sources"][0]["source"]))),
        "allocation: a Source placed twice is refused")
    check(any("not placed" in p for p in broken(
        lambda d: d["individual"][next(iter(d["individual"]))]["sources"].pop())),
        "allocation: an unplaced Source is refused")
    check(any("not in the current Brain" in p for p in broken(
        lambda d: d["training"].append("SRC-0999"))),
        "allocation: an unknown Source is refused")
    check(any("is titled" in p for p in broken(
        lambda d: d["training"].__setitem__(0, {"source": "SRC-0006", "title": "Wrong"}))),
        "allocation: a title that does not match the Brain is refused")
    check(any("at least two" in p for p in broken(
        lambda d: d["pairs"].__setitem__(next(iter(d["pairs"])),
                                         d["pairs"][next(iter(d["pairs"]))][:1]))),
        "allocation: a group must have at least two evaluators")
    check(not any("at least two" in p for p in broken(lambda d: None)),
          "allocation: a group of three is accepted")
    check(any("evaluator" in p for p in broken(
        lambda d: d["individual"][next(iter(d["individual"]))].__setitem__(
            "evaluator", "nobody"))),
        "allocation: an Individual split needs a listed evaluator")


# ================================================================= preflight
def test_preflight():
    report = preflight.check(BRAIN)
    check(report.ok, f"preflight: the current snapshot passes ({report.errors[:2]})")
    root = brain_copy()
    try:
        claims = root / "wiki" / "claims" / "claims.jsonl"
        rows = [json.loads(line) for line in claims.read_text().splitlines() if line]
        rows[0]["claim_object"] = "politics"
        rows[1]["concepts"] = ["CPT-does-not-exist"]
        rows[2]["datasets"] = ["DST-9999"]
        rows[3].pop("anchors")
        claims.write_text("\n".join(json.dumps(r) for r in rows) + "\n")
        edges = root / "wiki" / "graph" / "edges.jsonl"
        erows = [json.loads(line) for line in edges.read_text().splitlines() if line]
        erows[0]["grounding"] = "plausible"
        erows[1]["to"] = "CLM-9999-999"
        erows[2]["type"] = "COMPATIBLE_WITH"
        edges.write_text("\n".join(json.dumps(r) for r in erows) + "\n")
        (root / "schema" / "claim.md").write_text("changed")
        text = "\n".join(preflight.check(brain_module.load_brain_from(root)).errors)
        for fragment, label in (("claim_object 'politics'", "an invalid Claim object"),
                                ("CPT-does-not-exist does not resolve", "an unresolved Concept"),
                                ("DST-9999 does not resolve", "an unresolved Dataset"),
                                ("`anchors` is missing", "missing anchors"),
                                ("grounding 'plausible'", "an invalid grounding"),
                                ("do not resolve to Claims", "an unresolved endpoint"),
                                ("unknown type", "an obsolete relation type"),
                                ("schema/claim.md differs", "a changed schema file")):
            check(fragment in text, f"preflight: rejects {label}")
    finally:
        shutil.rmtree(root, ignore_errors=True)


# ================================================================= bootstrap
def test_bootstrap():
    book = fresh_round()
    meta = store.round_metadata()
    check(meta["round_id"] == ROUND and meta["eval_spec_version"] == "4.1"
          and meta["definitions_id"] == spec.definitions_id()
          and meta["brain_snapshot_id"] == BRAIN.snapshot_id
          and meta["brain_canonical_schema_version"] == "0.1.0",
          "bootstrap: ROUND records round, instrument, definitions and snapshot")
    check("RUN-2026-09-25-01" in meta["brain_runs"],
          "bootstrap: the runs and their schema versions are recorded")
    check(book.count(sheets.DEFINITIONS)
          == len(spec.definitions()["entries"]) + len(spec.CALIBRATION),
          "bootstrap: DEFINITIONS holds every definition and calibration shown")
    calibration = [r for r in book.read_tab(sheets.DEFINITIONS)
                   if r["key"].startswith("calibration.")]
    check(len(calibration) == len(spec.CALIBRATION)
          and all(r["source_file"] == "calibration meeting (spec 4.1)" for r in calibration)
          and not any(r["source_file"].startswith("calibration")
                      for r in book.read_tab(sheets.DEFINITIONS)
                      if not r["key"].startswith("calibration.")),
          "bootstrap: calibration rows keep their own provenance")
    check(sheets.RESTATEMENTS in book.tab_titles()
          and book.headers([sheets.RESTATEMENTS])[sheets.RESTATEMENTS]
          == ["restatement_key", "review_id", "source_id", "group_id", "claim_id",
              "active", "updated_at"],
          "bootstrap: a fresh round has the RESTATEMENTS tab")
    relation_header = book.headers([sheets.RELATION_RESPONSES])[sheets.RELATION_RESPONSES]
    check({"direction_answer", "direction_comment"} <= set(relation_header)
          and "related_claim_id" not in book.headers([sheets.RESPONSES])[sheets.RESPONSES],
          "bootstrap: Relations have Direction columns; RESPONSES has no related Claim")
    restatement_rows = [r for r in book.read_tab(sheets.RESPONSES)
                        if r["question_key"] == "SOURCE_RESTATEMENTS"]
    check(len(restatement_rows) == book.count(sheets.REVIEWS),
          "bootstrap: every review has its Restatements row")
    assignments = len(manifest.load_manifest_file()["assignments"])
    check(book.count(sheets.REVIEWS) == assignments
          and book.count(sheets.ASSIGNMENTS) == assignments,
          "bootstrap: one review per assignment")
    check(all(r["review_id"].startswith(ROUND + "|") for r in book.read_tab(sheets.REVIEWS)),
          "bootstrap: review ids carry the round namespace")
    check(bootstrap_round.verify_blocks(book, ROUND) == [],
          "bootstrap: every recorded block reads back at its rows")
    check(set(book.tab_titles()) == set(sheets.ALL_TABS), "bootstrap: exactly the round tabs")

    sys.argv = ["bootstrap_round.py"]
    with contextlib.redirect_stdout(io.StringIO()):
        check(bootstrap_round.main() == 1, "bootstrap: refuses a workbook holding a round")
    start(rid("agreement", EVALUATOR, AGR_SOURCE))
    sys.argv = ["bootstrap_round.py", "--overwrite-round", ROUND]
    with contextlib.redirect_stdout(io.StringIO()):
        check(bootstrap_round.main() == 1,
              "bootstrap: refuses to discard started evaluations without the flag")
    sys.argv += ["--discard-evaluations"]
    with contextlib.redirect_stdout(io.StringIO()):
        check(bootstrap_round.main() == 0, "bootstrap: an explicit overwrite replaces the round")

    other = fakes.FakeWorkbook()
    other.tabs["Notes"] = [["something"]]
    store.use_workbook(other)
    sys.argv = ["bootstrap_round.py"]
    with contextlib.redirect_stdout(io.StringIO()):
        check(bootstrap_round.main() == 1, "bootstrap: refuses a workbook with other content")


# ===================================================================== store
def test_store_roundtrip_and_conflicts():
    book = fresh_round()
    review = rid("individual", EVALUATOR, IND_SOURCE)
    ind_claim = BRAIN.claims_of(IND_SOURCE)[0]["id"]
    start(review)
    store.get_review(review)            # the page has already read REVIEWS by now
    before = dict(book.requests)
    data = store.load_review(review)
    check(book.requests["read"] - before.get("read", 0) == 1,
          "store: a paper's answers load in one read request")
    q2 = "CLAIM_Q02_CENTRAL_THESIS"
    row = data.response(q2, ind_claim)
    stamp = store.write(sheets.RESPONSES, row["response_key"], {**row, "answer": "No"},
                        rid=review, expected_updated_at="")
    check(store.load_review(review).response(q2, ind_claim)["answer"] == "No",
          "store: an answer round-trips")
    try:
        store.write(sheets.RESPONSES, row["response_key"], {**row, "answer": "Yes"},
                    rid=review, expected_updated_at="an older stamp")
        check(False, "store: a stale write is refused")
    except store.SaveConflict:
        check(True, "store: a stale write is refused")
    store.write(sheets.RESPONSES, row["response_key"], {**row, "answer": "Yes"},
                rid=review, expected_updated_at="an older stamp", also_accept={stamp})
    check(True, "store: the session's own earlier stamp is not a conflict")

    key = store.selection_key(review, ind_claim, "CPT-legal-drafting")
    values = {"review_id": review, "source_id": IND_SOURCE, "claim_id": ind_claim,
              "concept_id": "CPT-legal-drafting", "active": store.TRUE}
    store.upsert_selection(sheets.MISSING_CONCEPTS, key, values, rid=review)
    store.upsert_selection(sheets.MISSING_CONCEPTS, key, values, rid=review)
    check(book.count(sheets.MISSING_CONCEPTS) == 1, "store: a repeated selection is one row")
    store.upsert_selection(sheets.MISSING_CONCEPTS, key, {**values, "active": store.FALSE},
                           rid=review)
    check(store.load_review(review).active_missing(ind_claim) == [],
          "store: removal deactivates the same row")
    # Two sessions appending at the same instant: duplicated key in storage.
    book.append_rows(sheets.MISSING_CONCEPTS, [{**values, "selection_key": key,
                                                "active": store.TRUE,
                                                "updated_at": "9999-01-01"}])
    loaded = store.load_review(review)
    check(len(loaded.missing) == 1 and len(loaded.active_missing(ind_claim)) == 1,
          "store: duplicate keys collapse to one effective selection")
    try:
        store.write(sheets.RESPONSES, "no such key", {}, rid=review)
        check(False, "store: an unknown row is refused")
    except store.StorageIntegrityError:
        check(True, "store: an unknown row is refused")


def test_store_requires_a_workbook():
    store.use_workbook(None)
    try:
        store.workbook()
        check(False, "store: without a workbook nothing is stored")
    except sheets.StorageNotConfigured:
        check(True, "store: without a workbook nothing is stored")
    live = [f for f in APP.rglob("*.py") if "archive" not in f.parts and "tests" not in f.parts]
    check(not any(re.search(r"^\s*(import|from)\s+sqlite3", f.read_text(encoding="utf-8"), re.M)
                  for f in live), "store: no application module uses SQLite")


# ================================================================== progress
def test_progress_rules():
    fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    answer_everything(review, source)
    data = store.load_review(review)
    check(progress.missing_items(BRAIN, source, data) == [],
          "progress: a fully answered paper has nothing missing")
    claim = BRAIN.claims_of(source)[0]["id"]

    def with_response(key, object_id, **changes):
        copy = data.copy()
        copy.responses[store.response_lookup(key, object_id)].update(changes)
        return copy

    reasons = lambda d: [m.reason for m in progress.missing_items(BRAIN, source, d)]
    check(reasons(with_response("CLAIM_Q02_CENTRAL_THESIS", claim, answer="No")) == [],
          "progress: an optional comment never blocks completion")
    check(reasons(with_response("CLAIM_Q05_GROUNDING", claim, answer="In part")) == [],
          "progress: Question 5 comment is optional")
    dataset = BRAIN.dataset_ids_of(source)[0]
    check(any("comment required" in r for r in reasons(
        with_response("DATASET_DESCRIPTION", dataset, answer="No"))),
        "progress: the Description comment is required for No")
    check(any("text required" in r for r in reasons(
        with_response(spec.RECALL.key, source, answer=spec.RECALL_MOST))),
        "progress: Missing Claims is required below All")
    check(any("not evaluated" in r for r in reasons(
        with_response(spec.Q14.key, claim, answer=""))),
        "progress: Question 14 unanswered is not complete")
    check(any("not evaluated" in r for r in reasons(
        with_response(spec.Q14.key, claim, answer=spec.Q14_MISSING))),
        "progress: 'missing' with nothing selected is not complete")

    # dynamic totals: per assigned Concept, per candidate, per hosted Relation
    items = progress.claim_items(BRAIN, source, claim)
    check(sum(1 for i in items if i.question is spec.Q12)
          == len(BRAIN.concepts_of_claim(claim)), "progress: one Q12 per assigned Concept")
    check(sum(1 for i in items if i.question is spec.Q13)
          == len(BRAIN.candidates_of_claim(claim)), "progress: one Q13 per candidate only")
    check(sum(1 for i in items if i.question.unit == spec.UNIT_RELATION)
          == sum(3 if r["type"] in ("SUPPORTS", "ATTACKS") else 2
                 for r in BRAIN.relations_from_claim(claim)),
          "progress: three judgments per SUPPORTS/ATTACKS Relation, two per SAME_AS")
    check(not any(i.question.key == "CLAIM_Q01_RESTATEMENT" for i in items),
          "progress: no per-Claim restatement item")
    same_as = [r for s in BRAIN.sources for r in BRAIN.relations_hosted(s)
               if r["type"] == "SAME_AS"]
    if same_as:
        relation = same_as[0]
        keys = [i.question.key for i in progress.claim_items(
            BRAIN, relation["from_source"], relation["from"])
                if i.object_id == relation["key"]]
        check(keys == ["REL_GROUNDING", "REL_TYPE"],
              "progress: a SAME_AS relation has no Direction item")
    check(progress.dataset_items(BRAIN, "SRC-0003") == [],
          "progress: no Dataset question for a Source without Datasets")
    empty = store.ReviewData()
    check(all(i.question.unit != spec.UNIT_RELATION or
              any(r["key"] == i.object_id for r in BRAIN.relations_hosted(source))
              for i in progress.all_items(BRAIN, source, empty)),
          "progress: a Relation is evaluated only in its From Source")


def test_zero_claim_source():
    root = brain_copy()
    try:
        claims = root / "wiki" / "claims" / "claims.jsonl"
        rows = [r for r in (json.loads(l) for l in claims.read_text().splitlines() if l)
                if r["source"] != "SRC-0010"]
        claims.write_text("\n".join(json.dumps(r) for r in rows) + "\n")
        edges = root / "wiki" / "graph" / "edges.jsonl"
        kept = [e for e in (json.loads(l) for l in edges.read_text().splitlines() if l)
                if not str(e.get("from", "")).startswith("CLM-0010")
                and not str(e.get("to", "")).startswith("CLM-0010")]
        edges.write_text("\n".join(json.dumps(e) for e in kept) + "\n")
        copy = brain_module.load_brain_from(root)
        items = progress.all_items(copy, "SRC-0010", store.ReviewData())
        check([i.question.key for i in items] == [spec.RECALL.key, spec.RESTATEMENTS.key],
              "progress: a Source with no Claims is still evaluable by Claim recall")
    finally:
        shutil.rmtree(root, ignore_errors=True)


# ==================================================================== search
def test_concept_search():
    vocabulary = BRAIN.vocabulary(spec.concept_grid())
    results = conceptsearch.search("machine learning", vocabulary)
    check(results and results[0]["id"] == "CPT-machine-learning",
          "search: the exact name ranks first")
    check(conceptsearch.search("machine learning", vocabulary) == results,
          "search: deterministic")
    partial = conceptsearch.search("privacy xyzzy", vocabulary)
    check(any(e["id"] == "CPT-privacy-and-data-protection" for e in partial),
          "search: not every word has to match")
    check(len(conceptsearch.search("", vocabulary)) == len(vocabulary),
          "search: an empty query lists everything")
    entry = {"id": "CPT-alpha", "label": "Alpha", "family": "legal_task",
             "definition": "zeppelin"}
    check(conceptsearch.score("zeppelin", entry) == 0,
          "search: Concept definitions are neither search nor ranking evidence")
    check(conceptsearch.score("legal", entry) > 0, "search: the family is searchable")


# ================================================================== analysis
def test_agreement_metrics():
    pairs = [("Yes", "Yes"), ("No", "No"), ("Yes", "No"), ("Yes", "Yes")]
    kappa = analysis.cohen_kappa(pairs, ["Yes", "No"])
    check(abs(kappa - 0.5) < 1e-9, "analysis: Cohen's kappa on a known table")
    weighted = analysis.weighted_kappa([("No", "In part"), ("Yes", "Yes"), ("No", "No")],
                                       ["No", "In part", "Yes"])
    check(weighted is not None and 0 < weighted < 1, "analysis: linear-weighted kappa")
    book = fresh_round()
    group = sorted(a["evaluator_id"] for a in ASSIGNMENTS
                   if a["phase_id"] == "agreement" and a["source_id"] == AGR_SOURCE)
    for evaluator in group:
        review = rid("agreement", evaluator, AGR_SOURCE)
        start(review)
        answer_everything(review, AGR_SOURCE,
                          recall=spec.RECALL_ALL if evaluator == EVALUATOR else spec.RECALL_MOST)
    tables = analysis.export_tables()
    metrics = {m["question_key"]: m for m in tables["agreement_metrics"]}
    check(metrics["SOURCE_CLAIM_RECALL"]["metric"] == "linear_weighted_kappa"
          and metrics["CLAIM_Q05_GROUNDING"]["metric"] == "linear_weighted_kappa"
          and metrics["CLAIM_Q02_CENTRAL_THESIS"]["metric"] == "cohen_kappa",
          "analysis: the metric follows the question's scale")
    check("CLAIM_Q14_MISSING_CONCEPTS" not in metrics and "SOURCE_MISSING_CLAIMS" not in metrics,
          "analysis: structured and qualitative responses get no invented metric")
    check(metrics["REL_DIRECTION"]["metric"] == "cohen_kappa"
          and "SOURCE_RESTATEMENTS" not in metrics,
          "analysis: Direction is a binary judgment; restatements get no metric")
    check(all(r["eval_spec_version"] == "4.1" and r["round_id"] == ROUND
              for r in tables["claim_judgments"]),
          "analysis: every exported judgment names its round and instrument")
    unit_rows = [r for r in tables["agreement_pairs"]
                 if r["question_key"] == "SOURCE_CLAIM_RECALL"]
    expected_pairs = len(group) * (len(group) - 1) // 2
    check(len(unit_rows) == expected_pairs
          and {(r["evaluator_a"], r["evaluator_b"]) for r in unit_rows}
          == set(itertools.combinations(group, 2)),
          f"analysis: a group of {len(group)} yields every pair of its members per unit")
    check(len(tables["raw_responses"]) == book.count(sheets.RESPONSES),
          "analysis: the raw tabs are exported whole")


# ======================================================================== UI
def _app():
    from streamlit.testing.v1 import AppTest

    return AppTest.from_file(str(APP / "streamlit_app.py"), default_timeout=120)


def _login(at, name="Thiago"):
    at.run()
    at.text_input[0].input(os.environ["HE_APP_PASSWORD"])
    at.button[0].click()
    at.run()
    at.selectbox[0].select(name)
    at.run()
    at.button[0].click()
    at.run()


def _open(at, phase, source):
    at.sidebar.radio[0].set_value(manifest.PHASE_LABEL[phase])
    at.run()
    at.button(key=f"open|{phase}|{source}").click()
    at.run()


def _markdown(at) -> str:
    return "\n".join(m.value for m in at.markdown)


def _questions(at):
    """Answer controls on the page; the sidebar phase selector is not one."""
    return [r for r in at.radio if r.label != "Phase"]


def _flush(at):
    at.session_state["_save_queue"].flush(20)


def test_ui_guards():
    store.use_workbook(None)
    at = _app()
    at.run()
    check(at.title and "not available" in at.title[0].value,
          "ui: without a round workbook the app refuses to start")
    book = fresh_round()
    rows = book.tabs[sheets.ROUND]
    for row in rows:
        if row[0] == "round_id":
            row[1] = "ROUND-OTHER"
    store.forget_all()
    at = _app()
    at.run()
    check(at.title and "does not belong" in at.title[0].value,
          "ui: a workbook of another round is refused")


def test_ui_claim_flow():
    book = fresh_round()
    at = _app()
    _login(at)
    _open(at, "agreement", AGR_SOURCE)
    at.checkbox[0].check()
    at.run()
    [b for b in at.button if b.label == "Start evaluation"][0].click()
    at.run()
    page = _markdown(at)
    check("In tension with" not in page and "Attacked by" in page,
          "ui: the Source wiki shows current relation terminology")
    check(not _questions(at), "ui: the Source page asks nothing")
    at.button(key="nav|claims").click()
    at.run()
    check(not at.exception, f"ui: the Claim page renders ({[e.value for e in at.exception][:1]})")
    page = _markdown(at)
    check("<code>" not in page and "`" not in page,
          "ui: no schema name or value in code typography")
    check("Claim object" in page and "Law" in page and "Technology" in page,
          "ui: Claim object shows every category")
    check("Premise" not in page and "Positive form" not in page
          and "Plausibility" not in page and "Compatible with" not in page,
          "ui: removed fields and relation labels are absent")
    check("one per claim" in page, "ui: Question 2 shows the Statement definition")
    check("(this Claim)" in page and "From Claim" in page, "ui: Relation direction is shown")

    first = BRAIN.claims_of(AGR_SOURCE)[0]["id"]
    base = f"w|{rid('agreement', EVALUATOR, AGR_SOURCE)}"
    check("restatement of another Claim" not in page
          and not any(s.label == "Restatement of" for s in at.selectbox)
          and not any(r.key and "CLAIM_Q01" in r.key for r in at.radio),
          "ui: the Claim page has no restatement question or selector")

    # Question 14: search, add, clear search, add another; both persist.
    search_key = f"q14search|{rid('agreement', EVALUATOR, AGR_SOURCE)}|{first}"
    at.text_input(key=search_key).input("machine learning")
    at.run()
    at.button(key=f"{base}|q14add|{first}|CPT-machine-learning").click()
    at.run()
    at.text_input(key=search_key).input("")
    at.run()
    at.text_input(key=search_key).input("privacy")
    at.run()
    at.button(key=f"{base}|q14add|{first}|CPT-privacy-and-data-protection").click()
    at.run()
    at.text_input(key=search_key).input("")
    at.run()
    remove_keys = {b.key for b in at.button if b.label == "×"}
    page = _markdown(at)
    check({f"{base}|q14rm|{first}|CPT-machine-learning",
           f"{base}|q14rm|{first}|CPT-privacy-and-data-protection"} <= remove_keys
          and "Machine learning" in page and "Privacy and data protection" in page,
          "ui: selections persist when the search changes, each with its own ×")
    claim_concepts = set(BRAIN.concepts_of_claim(first))
    eligible = [e for e in BRAIN.vocabulary(spec.concept_grid())
                if e["id"] not in claim_concepts
                and e["id"] not in ("CPT-machine-learning", "CPT-privacy-and-data-protection")]
    offered = [b.key for b in at.button if b.key and f"|q14add|{first}|" in b.key]
    check(len(offered) == len(eligible),
          "ui: a cleared search offers the complete eligible vocabulary again")
    check(f"{base}|q14add|{first}|CPT-machine-learning" not in [b.key for b in at.button],
          "ui: a selected Concept is no longer offered")
    at.button(key=f"{base}|q14rm|{first}|CPT-machine-learning").click()
    at.run()
    check(f"{base}|q14rm|{first}|CPT-machine-learning" not in [b.key for b in at.button]
          and f"{base}|q14add|{first}|CPT-machine-learning" in [b.key for b in at.button],
          "ui: × removes a selected Concept and returns it to the browser")
    at.text_input(key=f"{base}|q14name|{first}").input("Legal chatbots")
    at.selectbox(key=f"{base}|q14family|{first}").select("technical_task")
    at.run()
    at.button(key=f"{base}|q14propose|{first}").click()
    at.run()
    check("Proposed · Legal chatbots" in _markdown(at)
          and any(b.key and "|q14rmp|" in b.key for b in at.button),
          "ui: a proposed Concept appears as selected, with its ×")
    _flush(at)
    stored = store.load_review(rid("agreement", EVALUATOR, AGR_SOURCE))
    check(stored.response(spec.Q14.key, first)["answer"] == spec.Q14_MISSING
          and [r["concept_id"] for r in stored.active_missing(first)]
          == ["CPT-privacy-and-data-protection"]
          and [r["name"] for r in stored.active_proposals(first)] == ["Legal chatbots"],
          "ui: Question 14 is stored as state, selections and proposals")
    check(not any(r["concept_id"] == "CPT-legal-chatbots" for r in
                  stored.missing.values()) and "CPT-legal-chatbots" not in BRAIN.concepts,
          "ui: a proposal creates no Brain Concept")
    check(book.count(sheets.MISSING_CONCEPTS) == 2,
          "ui: selecting and removing keeps one row per Concept")


def test_ui_no_datasets_and_completion():
    fresh_round()
    source = NO_DATASET_SOURCE
    review = rid(NO_DATASET_PHASE, EVALUATOR, source)
    start(review)
    answer_everything(review, source)
    at = _app()
    _login(at)
    _open(at, NO_DATASET_PHASE, source)
    at.button(key="nav|datasets").click()
    at.run()
    check(not _questions(at) and any("No Claim of this Source rests on a Dataset" in i.value
                               for i in at.info),
          "ui: no Dataset questions when the Source has no Dataset")
    at.button(key="nav|recall").click()
    at.run()
    check(not at.exception, "ui: Claim recall renders")
    at.button(key="nav|review").click()
    at.run()
    check(any("Every applicable item is answered" in s.value for s in at.success),
          "ui: Review reports a complete paper")
    page = _markdown(at)
    check("CITES" not in page and "Source attributes" not in page,
          "ui: Review holds only current evaluation items")
    [b for b in at.button if b.label == "Mark paper complete"][0].click()
    at.run()
    check(store.get_review(review)["status"] == store.STATUS_COMPLETE,
          "ui: a complete paper can be marked complete after a fresh re-read")


def test_ui_preflight_failure():
    fresh_round()
    root = brain_copy()
    try:
        claims = root / "wiki" / "claims" / "claims.jsonl"
        rows = [json.loads(line) for line in claims.read_text().splitlines() if line]
        rows[0]["concepts"] = ["CPT-nowhere"]
        claims.write_text("\n".join(json.dumps(r) for r in rows) + "\n")
        os.environ["HE_BRAIN_DIR"] = str(root)
        import streamlit as st
        st.cache_resource.clear()
        at = _app()
        at.run()
        check(at.title and at.title[0].value == "The evaluation is not available",
              "ui: an incompatible snapshot is rejected before evaluation begins")
        check(not at.text_input or at.text_input[0].label == "Admin secret",
              "ui: evaluators cannot sign in against an incompatible snapshot")
    finally:
        os.environ.pop("HE_BRAIN_DIR", None)
        import streamlit as st
        st.cache_resource.clear()
        shutil.rmtree(root, ignore_errors=True)


def test_ui_datasets_and_concepts():
    fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    at = _app()
    _login(at)
    _open(at, "agreement", source)
    base = f"w|{review}"
    claim = BRAIN.claims_of(source)[0]["id"]
    at.button(key="nav|claims").click()
    at.run()
    first, second = BRAIN.concepts_of_claim(claim)[:2]
    at.radio(key=f"{base}|{spec.Q12.key}|{claim}|{first}|a").set_value("No")
    at.run()
    check(at.radio(key=f"{base}|{spec.Q12.key}|{claim}|{second}|a").value is None,
          "ui: each Concept judgment is independent")
    check(any(t.label.startswith("Comment") for t in at.text_area),
          "ui: a compact comment appears under a Concept answered No")

    at.button(key="nav|datasets").click()
    at.run()
    dataset = BRAIN.dataset_ids_of(source)[0]
    description = f"{base}|DATASET_DESCRIPTION|{dataset}|a"
    at.radio(key=description).set_value("In part")
    at.run()
    check(any("A comment is required" in w.value for w in at.warning),
          "ui: Description In part asks for the required comment")
    check(any("describe what is missing or incorrect" in c.value for c in at.caption),
          "ui: the Description comment carries the agreed instruction")
    page = _markdown(at)
    check("Document type" not in page and "Agreement reported" not in page
          and "Used by" not in page, "ui: removed Dataset fields are absent")
    check("<code>" not in page and "`" not in page,
          "ui: no code typography on the Dataset page")
    at.button(key=f"ds_wiki|{dataset}").click()
    at.run()
    check(at.session_state["wiki_stack"] == [("dataset", dataset)],
          "ui: the Dataset wiki opens in the modal")
    at.session_state["wiki_stack"] = []
    at.run()
    check(at.radio(key=description).value == "In part",
          "ui: closing the modal keeps the Dataset answers")
    _flush(at)
    check(store.load_review(review).response("DATASET_DESCRIPTION", dataset)["answer"]
          == "In part", "ui: the Dataset answer is stored")


def test_ui_submission_locks():
    fresh_round()
    for source in manifest.assigned_source_ids("thiago", "training"):
        review = rid("training", "thiago", source)
        start(review)
        answer_everything(review, source)
        store.mark_complete(review)
    at = _app()
    _login(at)
    at.sidebar.radio[0].set_value("Training")
    at.run()
    at.checkbox(key="confirm_batch|training").check()
    at.run()
    [b for b in at.button if b.label == "Submit all papers"][0].click()
    at.run()
    check(store.phase_submitted(ROUND, "training", "thiago"),
          "ui: a complete phase can be finally submitted")
    check(all(store.get_review(rid("training", "thiago", s))["status"] == "submitted"
              for s in ("SRC-0006", "SRC-0009")), "ui: submission locks every paper")
    _open(at, "training", "SRC-0006")
    at.button(key="nav|claims").click()
    at.run()
    check(all(r.disabled for r in _questions(at)), "ui: a submitted paper is read-only")


def _expanders(at, label):
    return [e for e in at.expander if e.label == label]


def _inside(expander) -> str:
    return "\n".join(m.value for m in expander.markdown)


OBSOLETE_LABELS = ("Schema definition", "Concept definition", "Criterion definition",
                   "Extraction rule", "Authoritative extraction rule")


def _stored_rows(book) -> dict:
    return {tab: [list(r) for r in rows] for tab, rows in book.tabs.items()}


def test_ui_instructions_and_guidance():
    book = fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    at = _app()
    _login(at)
    _open(at, "agreement", source)
    at.button(key="nav|claims").click()
    at.run()
    instructions = _expanders(at, "Instructions")
    check(instructions and all(not e.proto.expanded for e in instructions),
          "ui: every Instructions expander starts closed")
    labels = [e.label for e in at.expander] + [p.label for p in at.get("popover")]
    headings = [m.value for m in at.markdown if m.value.lstrip().startswith("#")]
    check(not any(old in label for label in labels + headings for old in OBSOLETE_LABELS),
          "ui: no obsolete guidance label on any container or heading")
    everything = [_inside(e) for e in instructions]

    def plain(key, field="text"):
        return html.escape(spec.plain(spec.definition(key)[field]))

    def holds(key) -> str:
        lines = [html.escape(line) for line in spec.definition(key)["text"].split("\n")
                 if line.strip()]
        return next((t for t in everything if all(line in t for line in lines)), "")

    skill = freeze_definitions.DEFAULT_SKILLS / "extract-claims" / "SKILL.md"
    for key, question in (("skill.extract_claims.modality", "CLAIM_Q03_MODALITY"),
                          ("skill.extract_claims.standalone", "CLAIM_Q04_STANDALONE")):
        stored = spec.definition(key)["text"]
        if skill.exists():
            check(stored in skill.read_text(encoding="utf-8"),
                  f"ui: {key} is frozen exactly as the extract-claims skill words it")
        check(holds(key), f"ui: {question} shows its extract-claims passage in Instructions")
        check(spec.BY_KEY[question].text == AGREED[question],
              f"ui: {question} wording is unchanged")
    check('"may"' in spec.definition("skill.extract_claims.modality")["text"]
          and "`" not in spec.definition("skill.extract_claims.modality")["text"],
          "ui: the frozen passage is stored untransformed")

    q2 = next((t for t in everything if plain("claim.node") in t), "")
    check(plain("claim.statement") in q2 and "**Claim**" in q2 and "**Statement**" in q2,
          "ui: Question 2's Instructions hold the Claim and Statement passages")
    claim_object = next((t for t in everything if plain("claim.claim_object") in t), "")
    check(all(html.escape(spec.plain(i["text"])) in claim_object
              for i in spec.definition("claim.claim_object")["items"])
          and all(v in claim_object for v in ("Law", "Technology", "Other")),
          "ui: Claim object's Instructions hold all three values and meanings")
    check("(assigned)" not in "\n".join(everything),
          "ui: nothing is added inside a frozen definition")
    grounding = spec.plain(spec.definition("edge.grounding")["verbatim"])
    check(any(grounding in t for t in everything),
          "ui: Grounding Instructions hold the create-edges rule word for word")
    page = _markdown(at)
    check(all(term in page and html.escape(text) in page
              for term, text in spec.GROUNDING_QUICK_GUIDE)
          and "Quick guide" in page, "ui: the Grounding Quick guide is visible")
    check(not any(grounding in m.value for m in at.markdown if m.value in page
                  and "Quick guide" in m.value),
          "ui: the Quick guide does not replace the authoritative rule")

    claim = BRAIN.claims_of(source)[0]
    check(html.escape(claim["anchors"][0]["quote"]) in page and "Assigned value" in page,
          "ui: anchors and assigned values stay outside the Instructions")
    shown = page + "\n".join(c.value for c in at.caption)
    definitions = [BRAIN.concept(c)["definition"] for c in BRAIN.concepts_of_claim(claim["id"])]
    eligible = [e for e in BRAIN.vocabulary(spec.concept_grid()) if e["definition"]]
    check(not any(html.escape(d) in shown or d in shown for d in definitions),
          "ui: Q12 and Q13 show no Concept definition")
    check(not any(e["definition"] in shown or html.escape(e["definition"]) in shown
                  for e in eligible), "ui: the Q14 browser shows no Concept definition")
    first = BRAIN.concepts_of_claim(claim["id"])[0]
    family = brain_module.family_label(BRAIN.concept(first)["concept_type"])
    check(any(c.value.startswith(family) for c in at.caption),
          "ui: Concept family is shown with each Concept")
    concept_captions = [c.value for c in at.caption]
    for concept_id in BRAIN.concepts_of_claim(claim["id"]):
        concept = BRAIN.concept(concept_id)
        check(brain_module.family_label(concept["concept_type"]) in concept_captions,
              f"ui: {concept_id} shows its family")
    statuses = (" · Anchor", " · Candidate", " · Emergent")
    q12 = [c for c in concept_captions
           if c in {brain_module.family_label(f) for f in brain_module.CONCEPT_FAMILIES}
           or any(c.endswith(s) for s in statuses)]
    candidates = BRAIN.candidates_of_claim(claim["id"])
    check(sum(1 for c in q12 if any(c.endswith(s) for s in statuses)) == len(candidates)
          and all(c.endswith(" · Candidate") for c in q12
                  if any(c.endswith(s) for s in statuses)),
          "ui: Q12 and the Q14 browser show no status; only Q13 says Candidate")
    check(spec.PROPOSAL_HINT in concept_captions,
          "ui: the new-Concept hint is visible above the proposal form")

    panel = [m.value for m in at.markdown if "Claim under review" in m.value]
    check(panel and claim["id"] in panel[0] and html.escape(claim["statement"]) in panel[0],
          "ui: the Claim panel shows the current Claim id and statement")
    css = "\n".join(m.value for m in at.markdown if "<style>" in m.value)
    check(".st-key-claim_sticky" in css and 'data-testid="stLayoutWrapper"' in css
          and "position:sticky" in css, "ui: the Claim panel carries its sticky rule")
    before = _stored_rows(book)
    at.button(key="next_claim").click()
    at.run()
    second = BRAIN.claims_of(source)[1]
    panel = [m.value for m in at.markdown if "Claim under review" in m.value]
    check(panel and second["id"] in panel[0] and html.escape(second["statement"]) in panel[0],
          "ui: the Claim panel follows Claim navigation")
    check(_stored_rows(book) == before, "ui: the panel and navigation write nothing")

    at.button(key="nav|datasets").click()
    at.run()
    instructions = _expanders(at, "Instructions")
    check(instructions and all(not e.proto.expanded for e in instructions),
          "ui: Dataset Instructions start closed")
    jurisdiction = next((_inside(e) for e in instructions
                         if plain("dataset.jurisdiction") in _inside(e)), "")
    check("**Dataset jurisdiction**" in jurisdiction
          and "**Referenced jurisdiction definition**" in jurisdiction
          and plain("claim.claim_jurisdiction") in jurisdiction,
          "ui: Dataset Jurisdiction shows two separately headed passages")
    check(spec.definition("dataset.jurisdiction")["text"]
          == "legal system(s) the texts come from; jurisdiction codes in `schema/claim.md`.",
          "ui: the stored passage carries no interface heading")
    check(len(_questions(at)) == 6, "ui: the Dataset page keeps all six questions")


def test_ui_raw_markdown():
    book = fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    at = _app()
    _login(at)
    _open(at, "agreement", source)
    before = _stored_rows(book)
    cached = at.session_state["_review_cache"][review].copy()
    raw = [e for e in at.expander if e.label == "Raw Markdown"]
    on_disk = (BRAIN.root / "wiki" / "sources" / f"{source}.md").read_bytes().decode("utf-8")
    check(raw and any(c.value == on_disk for c in raw[0].code),
          "ui: Source Raw Markdown is the exact file on disk, frontmatter included")
    check(on_disk.startswith("---") and on_disk != BRAIN.source_body(source),
          "ui: the raw view is the file itself, not the rendered body")
    at.button(key="nav|datasets").click()
    at.run()
    dataset = BRAIN.dataset_ids_of(source)[0]
    at.button(key=f"ds_wiki|{dataset}").click()
    at.run()
    dataset_raw = (BRAIN.root / "wiki" / "datasets" / f"{dataset}.md").read_bytes().decode("utf-8")
    check(any(c.value == dataset_raw for c in at.get("code")),
          "ui: Dataset wiki Raw Markdown in the modal is the exact file")
    check(_stored_rows(book) == before
          and at.session_state["_review_cache"][review].responses == cached.responses,
          "ui: opening Raw Markdown changes no stored or cached evaluation data")


def test_ui_review_groups_missing_items():
    fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    at = _app()
    _login(at)
    _open(at, "agreement", source)
    at.button(key="nav|review").click()
    at.run()
    labels = [e.label for e in at.expander]
    claim_groups = [l for l in labels if l.startswith("Claim CLM-") and "missing" in l]
    dataset_groups = [l for l in labels if l.startswith("Dataset DST-") and "missing" in l]
    check(len(claim_groups) == len(BRAIN.claims_of(source)),
          "ui: missing items are grouped by Claim")
    check(len(dataset_groups) == len(BRAIN.dataset_ids_of(source)),
          "ui: missing Dataset items are grouped by Dataset")
    check(any(b.key == "goto|recall|0" for b in at.button),
          "ui: Claim recall is listed separately")
    check(not any(b.key and re.fullmatch(r"goto\|\d+", b.key) for b in at.button),
          "ui: no flat source-wide list of individual missing items")
    check(all(not e.proto.expanded for e in at.expander if "missing" in e.label),
          "ui: each group starts closed")


def test_ui_group_of_three():
    fresh_round()
    sizes = {e: manifest.agreement_group_size(e) for e in ("thiago", "giovanni")}
    check(sizes["thiago"] == 3 and sizes["giovanni"] == 2,
          "ui: Agreement group sizes come from the assignment rows (3 and 2)")
    at = _app()
    _login(at)
    at.sidebar.radio[0].set_value("Agreement")
    at.run()
    captions = " ".join(c.value for c in at.caption)
    check("You and 2 other evaluators assess the same papers" in captions
          and "one other evaluator" not in captions,
          "ui: a member of a group of three is told there are two other evaluators")
    check(not any(n in captions for n in ("Francesca", "Thibault")),
          "ui: group-mates are not named")
    at.sidebar.radio[0].set_value("Admin")
    at.run()
    at.text_input[0].input(os.environ["HE_ADMIN_SECRET"])
    at.button[0].click()
    at.run()
    table = at.dataframe[0].value
    check("group" in table.columns, "ui: Admin progress has a group column")
    paper = sorted(a["source_id"] for a in ASSIGNMENTS
                   if a["phase_id"] == "agreement" and a["evaluator_id"] == "thiago")[0]
    members = table[(table["phase"] == "agreement") & (table["paper"] == paper)]
    check(len(members) == 3 and set(members["group"]) == {"P1"},
          "ui: Admin progress lists all three members of P1 for one Agreement paper")


# ================================================================ 4.1: store
def test_restatement_storage_and_rules():
    book = fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    claims = [c["id"] for c in BRAIN.claims_of(source)]
    a, b, c, d = claims[:4]
    check(store.restatement_group_id([c, a, b]) == store.restatement_group_id([b, c, a, a])
          and store.restatement_group_id([a, b]) != store.restatement_group_id([a, c])
          and store.restatement_group_id([a, b]).startswith("RG-"),
          "restatements: the group id depends only on the set of Claims")

    data = store.load_review(review)
    item = next(i for i in progress.recall_items(BRAIN, source, data)
                if i.question is spec.RESTATEMENTS)
    problem = lambda d: progress.item_problem(item, d)
    check(problem(data) == "unanswered", "restatements: required, unanswered at first")

    def respond(d, answer):
        d = d.copy()
        d.responses[store.response_lookup(spec.RESTATEMENTS.key, source)]["answer"] = answer
        return d

    def with_group(d, ids, active=True):
        d = d.copy()
        gid = store.restatement_group_id(ids)
        for cid in ids:
            d.restatements[store.pair_lookup(gid, cid)] = {
                "group_id": gid, "claim_id": cid,
                "active": store.TRUE if active else store.FALSE}
        return d

    check(problem(respond(data, "none")) is None, "restatements: No with no group completes")
    check(problem(respond(data, "present")) is not None,
          "restatements: Yes without a group is not complete")
    one = with_group(respond(data, "present"), [a, b])
    check(problem(one) is None, "restatements: Yes with one valid group completes")
    two = with_group(one, [c, d])
    check(problem(two) is None and len(two.active_restatement_groups()) == 2,
          "restatements: several groups can coexist")
    check(problem(with_group(respond(data, "none"), [a, b])) is not None,
          "restatements: No with an active group is not complete")
    check(problem(with_group(respond(data, "none"), [a, b], active=False)) is None,
          "restatements: removed groups do not count")
    check(progress.item_problem(item, with_group(data, [a, b])) == "unanswered",
          "restatements: groups never imply the answer Yes")

    # data-boundary validation
    ok = lambda ids, d=data: progress.validate_restatement_group(BRAIN, source, d, ids)
    check(ok([b, a, a]) == (sorted([a, b]), None), "restatements: duplicates are dropped")
    check(ok([a, a])[1] is not None, "restatements: one Claim twice is not a group")
    check(ok([a])[1] is not None, "restatements: a group needs two Claims")
    other = next(x["id"] for s in BRAIN.sources if s != source for x in BRAIN.claims_of(s))
    check(ok([a, other])[1] is not None,
          "restatements: a Claim of another Source is refused")
    check("already assigned to another restatement group" in (ok([a, c], one)[1] or ""),
          "restatements: a Claim cannot join a second active group")

    # malformed stored groups are ignored for completion and reported
    bad = respond(data, "present")
    bad.restatements["x|" + a] = {"group_id": store.restatement_group_id([a, b]),
                                  "claim_id": a, "active": store.TRUE}
    check(problem(bad) is not None and "invalid" in problem(bad),
          "restatements: a one-member stored group is flagged")
    overlap = with_group(with_group(respond(data, "present"), [a, b]), [a, c])
    check(problem(overlap) is not None,
          "restatements: a Claim in two stored groups is flagged")
    foreign = respond(data, "present")
    gid = store.restatement_group_id([a, other])
    for cid in (a, other):
        foreign.restatements[store.pair_lookup(gid, cid)] = {
            "group_id": gid, "claim_id": cid, "active": store.TRUE}
    check(problem(foreign) is not None,
          "restatements: a stored group with another Source's Claim is flagged")

    # storage round trip, idempotence, removal, reactivation
    gid = store.restatement_group_id([a, b])
    rows = {store.restatement_key(review, gid, cid): {
        "review_id": review, "source_id": source, "group_id": gid, "claim_id": cid,
        "active": store.TRUE} for cid in (a, b)}
    store.upsert_selections(sheets.RESTATEMENTS, rows, rid=review)
    store.upsert_selections(sheets.RESTATEMENTS, rows, rid=review)
    check(book.count(sheets.RESTATEMENTS) == 2,
          "restatements: a group of two is two membership rows, written once")
    check(store.load_review(review).active_restatement_groups() == {gid: sorted([a, b])},
          "restatements: groups survive a reload")
    store.upsert_selections(sheets.RESTATEMENTS, {k: {**v, "active": store.FALSE}
                                                   for k, v in rows.items()}, rid=review)
    check(store.load_review(review).active_restatement_groups() == {}
          and book.count(sheets.RESTATEMENTS) == 2
          and all(r["active"] == store.FALSE for r in book.read_tab(sheets.RESTATEMENTS)),
          "restatements: removal marks the same rows inactive")
    store.upsert_selections(sheets.RESTATEMENTS, rows, rid=review)
    check(book.count(sheets.RESTATEMENTS) == 2
          and store.load_review(review).active_restatement_groups() == {gid: sorted([a, b])},
          "restatements: re-adding the same set reactivates the same rows")
    tables = analysis.export_tables()
    check([(r["group_id"], r["claim_id"]) for r in tables["restatement_groups"]]
          == [(gid, a), (gid, b)] and tables["restatement_groups"][0]["round_id"] == ROUND
          and len(tables["raw_restatements"]) == 2,
          "restatements: exported as one row per Claim membership, with provenance")
    check(analysis.sheet_names(tables) == {n: n for n in tables}
          and len(set(analysis.sheet_names(["x" * 40, "x" * 35]).values())) == 2,
          "exports: XLSX tabs are named after their tables, unique within 31 characters")


def test_export_xlsx():
    import openpyxl

    fresh_round()
    review = rid("agreement", EVALUATOR, AGR_SOURCE)
    start(review)
    answer_everything(review, AGR_SOURCE)
    row = store.load_review(review).response("CLAIM_Q02_CENTRAL_THESIS",
                                              BRAIN.claims_of(AGR_SOURCE)[0]["id"])
    store.write(sheets.RESPONSES, row["response_key"],
                {**row, "answer": "No", "comment": "pasted\x0btext\x01 here"}, rid=review)
    tables = analysis.export_tables()
    tables["restatement_groups"] = []
    book = openpyxl.load_workbook(io.BytesIO(analysis.export_xlsx(tables)))
    check(book.sheetnames == list(tables),
          "exports: the XLSX has one tab per table, named after it, in order")
    check(all(book[name].max_row - 1 == len(rows) for name, rows in tables.items() if rows),
          "exports: every tab holds all of its table's rows")
    comments = [cell.value for r in book["claim_judgments"].iter_rows() for cell in r]
    check("pastedtext here" in comments,
          "exports: characters Excel cannot store are dropped, the rest kept")
    check("restatement_groups" in book.sheetnames
          and book["restatement_groups"].max_row <= 1,
          "exports: an empty table still has its tab")
    raw = book["raw_restatements"]
    check([c.value for c in raw[1]] == list(sheets.COLUMNS[sheets.RESTATEMENTS]),
          "exports: an empty raw tab keeps its header")


def test_ui_admin_exports():
    book = fresh_round()
    at = _app()
    _login(at)
    at.sidebar.radio[0].set_value("Admin")
    at.run()
    at.text_input[0].input(os.environ["HE_ADMIN_SECRET"])
    at.button[0].click()
    at.run()
    before = _stored_rows(book)
    at.button(key="prepare_exports").click()
    at.run()
    downloads = at.get("download_button")
    labels = [d.proto.label for d in downloads]
    check(labels and labels[0] == "Download all tables (.xlsx)",
          "ui: Exports offers the single XLSX first")
    csv = [e for e in at.expander if e.label == "Individual tables (CSV)"]
    check(csv and not csv[0].proto.expanded
          and len(csv[0].get("download_button")) == len(analysis.export_tables()),
          "ui: the per-table CSVs are in a closed expander")
    check(_stored_rows(book) == before, "ui: preparing exports writes nothing")


# ============================================================ save semantics
def test_save_conflicts_are_per_record():
    import savequeue

    book = fresh_round()
    review = rid("individual", EVALUATOR, IND_SOURCE)
    start(review)
    rows = list(store.load_review(review).responses.values())[:3]
    stale = store.write(sheets.RESPONSES, rows[0]["response_key"],
                        {**rows[0], "answer": "No"}, rid=review)
    records = [{"tab": sheets.RESPONSES, "key": r["response_key"],
                "values": {**r, "answer": "Yes"}, "rid": review,
                "expected_updated_at": "an older stamp" if i == 0 else ""}
               for i, r in enumerate(rows)]
    results = store.save_many(records)
    stored = store.load_review(review)
    answer = lambda r: stored.response(r["question_key"], r["object_id"])["answer"]
    check(isinstance(results[0], store.SaveConflict) and answer(rows[0]) == "No",
          "save: a stale record is refused and the stored value stands")
    check(all(isinstance(x, str) for x in results[1:])
          and [answer(r) for r in rows[1:]] == ["Yes", "Yes"],
          "save: the other records of the same batch are still written")

    queue = savequeue.SaveQueue()
    queue.set_batch_writer(lambda f, a, k: {"tab": a[0], "key": a[1], "values": a[2], **k}
                           if f is store.write else None, store.save_many)
    queue.set_fatal(store.SaveConflict)
    fresh = store.load_review(review)
    current = [fresh.responses[store.response_lookup(r["question_key"], r["object_id"])]
               for r in rows]
    with queue._lock:                   # both queued before the worker takes any
        for i, row in enumerate(current):
            queue._pending[row["response_key"]] = (
                store.write, (sheets.RESPONSES, row["response_key"],
                              {**row, "answer": "In part" if i == 0 else "No"}),
                {"rid": review, "expected_updated_at": "stale" if i == 0 else
                 row["updated_at"], "also_accept": set()}, None)
    queue._ensure_worker()
    queue.flush(10)
    stored = store.load_review(review)
    check(list(queue.conflicts) == [current[0]["response_key"]]
          and [answer(r) for r in rows[1:]] == ["No", "No"],
          "save queue: only the conflicting answer is reported; the others are stored")


def test_save_retry_after_lost_response():
    import savequeue

    book = fresh_round()
    review = rid("individual", EVALUATOR, IND_SOURCE)
    start(review)
    row = next(iter(store.load_review(review).responses.values()))
    first = store.write(sheets.RESPONSES, row["response_key"], {**row, "answer": "Yes"},
                        rid=review)
    real, calls = book.update_rows, {"n": 0}

    def lands_then_times_out(tab, updates):
        real(tab, updates)
        if tab == sheets.RESPONSES:
            calls["n"] += 1
            if calls["n"] == 1:
                raise sheets.StorageError("the response was lost after the write landed")

    book.update_rows = lands_then_times_out
    old_delay, savequeue.RETRY_SECONDS = savequeue.RETRY_SECONDS, 0.01
    try:
        queue = savequeue.SaveQueue()
        queue.set_fatal(store.SaveConflict)
        queue.submit("k", store.write, sheets.RESPONSES, row["response_key"],
                     {**row, "answer": "No", "updated_at": first}, rid=review,
                     expected_updated_at=first, also_accept={first})
        queue.flush(10)
    finally:
        savequeue.RETRY_SECONDS = old_delay
        book.update_rows = real
    stored = store.load_review(review).response(row["question_key"], row["object_id"])
    check(not queue.conflicts and not queue.failures and stored["answer"] == "No",
          "save: a retry whose first attempt landed is not a conflict")
    elsewhere = store.write(sheets.RESPONSES, row["response_key"],
                            {**stored, "answer": "Yes", "comment": "other tab"}, rid=review)
    try:
        store.write(sheets.RESPONSES, row["response_key"], {**stored, "answer": "No"},
                    rid=review, expected_updated_at=stored["updated_at"])
        check(False, "save: a real concurrent edit is still refused")
    except store.SaveConflict:
        check(elsewhere and True, "save: a real concurrent edit is still refused")


# ================================================================ integrity
def _audit(book):
    import integrity

    return integrity.run(BRAIN, book, ROUND)


def test_integrity_audit():
    import integrity

    book = fresh_round()
    levels = lambda ps, level: [p for p in ps if p.level == level]
    check(not levels(_audit(book), integrity.ERROR)
          and not levels(_audit(book), integrity.WARNING),
          "integrity: a freshly bootstrapped round is clean")
    review = rid("agreement", EVALUATOR, AGR_SOURCE)
    start(review)
    answer_everything(review, AGR_SOURCE)
    store.mark_complete(review)
    check(not levels(_audit(book), integrity.ERROR),
          "integrity: a completed, fully answered paper is clean")

    def row_of(tab, predicate):
        header = list(sheets.COLUMNS[tab])
        for raw in book.tabs[tab][1:]:
            record = dict(zip(header, raw))
            if predicate(record):
                return raw, header
        raise AssertionError(tab)

    def reported(fragment):
        return any(fragment in p.what for p in levels(_audit(book), integrity.ERROR))

    claim = BRAIN.claims_of(AGR_SOURCE)[0]["id"]
    raw, header = row_of(sheets.RESPONSES, lambda r: r["review_id"] == review
                         and r["question_key"] == "CLAIM_Q02_CENTRAL_THESIS"
                         and r["object_id"] == claim)
    raw[header.index("answer")] = "Maybe"
    check(reported("is not one of its options"), "integrity: an invalid answer is an error")
    raw[header.index("answer")] = ""
    check(reported("marked complete but"),
          "integrity: a complete paper with an item missing is an error")
    raw[header.index("answer")] = "Yes"

    raw, header = row_of(sheets.RESPONSES, lambda r: r["review_id"] == review
                         and r["question_key"] == spec.Q14.key and r["object_id"] == claim)
    raw[header.index("answer")] = spec.Q14_MISSING
    check(any("none is selected" in p.what for p in _audit(book)),
          "integrity: Question 14 'missing' without a selection is reported")
    raw[header.index("answer")] = spec.Q14_NONE_MISSING

    raw, header = row_of(sheets.RESPONSES, lambda r: r["review_id"] == review
                         and r["question_key"] == spec.RESTATEMENTS.key)
    raw[header.index("answer")] = spec.RESTATEMENTS_PRESENT
    gid = store.restatement_group_id([claim])
    book.append_rows(sheets.RESTATEMENTS, [{
        "restatement_key": store.restatement_key(review, gid, claim),
        "review_id": review, "source_id": AGR_SOURCE, "group_id": gid,
        "claim_id": claim, "active": store.TRUE, "updated_at": "x"}])
    check(reported("fewer than two Claims"),
          "integrity: a one-member restatement group is an error")
    book.tabs[sheets.RESTATEMENTS] = book.tabs[sheets.RESTATEMENTS][:1]
    raw[header.index("answer")] = spec.RESTATEMENTS_NONE

    same_as = [r for r in book.read_tab(sheets.RELATION_RESPONSES)
               if r["relation_type"] == "SAME_AS"]
    if same_as:
        raw, header = row_of(sheets.RELATION_RESPONSES,
                             lambda r: r["response_key"] == same_as[0]["response_key"])
        raw[header.index("direction_answer")] = "Yes"
        check(reported("does not ask it"),
              "integrity: a Direction answer on SAME_AS is an error")
        raw[header.index("direction_answer")] = ""

    raw, header = row_of(sheets.REVIEWS, lambda r: r["review_id"] == review)
    raw[header.index("status")] = store.STATUS_SUBMITTED
    check(reported("no submission marker"),
          "integrity: a submitted paper without its phase marker is an error")
    raw[header.index("status")] = store.STATUS_COMPLETE
    first = raw[header.index("responses_first_row")]
    raw[header.index("responses_first_row")] = str(int(first) + 1)
    check(reported("are not all this review's") or reported("its block has"),
          "integrity: a shifted block range is an error")
    raw[header.index("responses_first_row")] = first
    check(not levels(_audit(book), integrity.ERROR),
          "integrity: undoing each corruption leaves the round clean")


# ================================================================ 4.1: UI
def _hosting(types) -> tuple[str, str, int, dict] | None:
    """(phase, source, claim index, relation) for Thiago's first relation of a type."""
    for phase in ("training", "agreement", "individual"):
        for source in _assigned(EVALUATOR, phase):
            for index, claim in enumerate(BRAIN.claims_of(source)):
                for relation in BRAIN.relations_from_claim(claim["id"]):
                    if relation["type"] in types:
                        return phase, source, index, relation
    return None


def _open_claim(at, phase, source, index):
    _open(at, phase, source)
    at.session_state[f"claim_idx|{rid(phase, EVALUATOR, source)}"] = index
    at.button(key="nav|claims").click()
    at.run()


def test_ui_relation_direction():
    for types, label in ((("SUPPORTS",), "SUPPORTS"), (("ATTACKS",), "ATTACKS")):
        found = _hosting(types)
        check(found is not None, f"ui: an assigned {label} relation exists for the check")
        if not found:
            continue
        phase, source, index, relation = found
        book = fresh_round()
        review = rid(phase, EVALUATOR, source)
        start(review)
        at = _app()
        _login(at)
        _open_claim(at, phase, source, index)
        base = f"w|{review}"
        keys = {q: f"{base}|{q}|{relation['key']}|a"
                for q in ("REL_GROUNDING", "REL_DIRECTION", "REL_TYPE")}
        radios = {r.key: r for r in at.radio}
        check(all(k in radios for k in keys.values()),
              f"ui: {label} shows Grounding, Direction and Relation type")
        check(radios[keys["REL_DIRECTION"]].options == ["Yes", "No"],
              f"ui: {label} Direction is Yes / No")
        page = _markdown(at)
        check(f"{relation['from']} → {relation['to']}" in page
              and spec.DIRECTION.text in page, f"ui: {label} shows the arrow and the question")
        direction = [_inside(e) for e in _expanders(at, "Instructions")
                     if spec.CALIBRATION["relation.direction"][0] in _inside(e)]
        check(direction and "Schema / skill" in direction[0] and "Calibration" in direction[0]
              and html.escape(spec.plain(spec.definition("edge.from_to")["text"]))
              in direction[0],
              f"ui: {label} Direction Instructions hold from/to and the calibration apart")
        at.radio(key=keys["REL_DIRECTION"]).set_value("No")
        at.run()
        type_radio = at.radio(key=keys["REL_TYPE"])
        check(not type_radio.disabled and type_radio.value is None
              and at.radio(key=keys["REL_GROUNDING"]).value is None,
              f"ui: {label} Direction No leaves Grounding and Type open and unanswered")
        check(any(t.label.startswith("Correction or comment") for t in at.text_area),
              f"ui: {label} Direction No offers an optional comment")
        at.radio(key=keys["REL_TYPE"]).set_value("Yes")
        at.run()
        _flush(at)
        stored = store.load_review(review).relation(relation["key"])
        check(stored["direction_answer"] == "No" and stored["type_answer"] == "Yes"
              and stored["grounding_answer"] == "",
              f"ui: {label} judgments are stored independently")
        at.button(key="nav|review").click()
        at.run()
        page = _markdown(at)
        check("Direction · " in page and "Relation type · " in page and "Grounding · " in page,
              f"ui: Review lists {label} Grounding, Direction and Type separately")
        tables = analysis.export_tables()
        check(any(r["question_key"] == "REL_DIRECTION" and r["relation_key"] == relation["key"]
                  and r["answer"] == "No" for r in tables["relation_judgments"]),
              f"ui: the {label} Direction judgment is exported")

    found = _hosting(("SAME_AS",))
    check(found is not None, "ui: an assigned SAME_AS relation exists for the check")
    if found:
        phase, source, index, relation = found
        fresh_round()
        review = rid(phase, EVALUATOR, source)
        start(review)
        at = _app()
        _login(at)
        _open_claim(at, phase, source, index)
        keys = {r.key for r in at.radio}
        base = f"w|{review}"
        check(f"{base}|REL_GROUNDING|{relation['key']}|a" in keys
              and f"{base}|REL_TYPE|{relation['key']}|a" in keys
              and f"{base}|REL_DIRECTION|{relation['key']}|a" not in keys,
              "ui: SAME_AS keeps Grounding and Type and asks no Direction")
        tables = analysis.export_tables()
        check(not any(r["relation_key"] == relation["key"] and r["question_key"] == "REL_DIRECTION"
                      for r in tables["relation_judgments"]),
              "ui: no Direction row is exported for SAME_AS")


def test_ui_restatements():
    book = fresh_round()
    source = AGR_SOURCE
    review = rid("agreement", EVALUATOR, source)
    start(review)
    at = _app()
    _login(at)
    _open(at, "agreement", source)
    at.button(key="nav|recall").click()
    at.run()
    base = f"w|{review}"
    answer = f"{base}|SOURCE_RESTATEMENTS|{source}|a"
    pick = f"{base}|rg_pick|{source}"
    add = f"{base}|rg_add|{source}"
    page = _markdown(at)
    order = [page.find(t) for t in ("Extracted Claims", spec.RECALL.text,
                                    "#### Restatements")]
    check(all(i >= 0 for i in order) and order == sorted(order),
          "ui: Claim recall shows Extracted Claims, completeness, then Restatements")
    radio = at.radio(key=answer)
    check(radio.options == ["No", "Yes"] and radio.value is None and not radio.disabled,
          "ui: Restatements is an unanswered No / Yes")
    radio.set_value("none")
    at.run()
    _flush(at)
    check(store.load_review(review).response("SOURCE_RESTATEMENTS", source)["answer"] == "none",
          "ui: No is stored as none")
    at.radio(key=answer).set_value("present")
    at.run()
    check(any(m.key == pick for m in at.multiselect) and at.button(key=add).disabled,
          "ui: Yes offers the Claim selector; Add needs two Claims")
    claims = [c["id"] for c in BRAIN.claims_of(source)]
    a, b, c, d = claims[:4]
    at.multiselect(key=pick).select(b)
    at.run()
    check(at.button(key=add).disabled, "ui: one Claim cannot make a group")
    at.multiselect(key=pick).select(a)
    at.run()
    at.button(key=add).click()
    at.run()
    check(at.radio(key=answer).disabled
          and any("Remove the restatement groups to answer No." in x.value for x in at.caption),
          "ui: with a group, No is unavailable and the reason is given")
    at.multiselect(key=pick).select(a)
    at.multiselect(key=pick).select(c)
    at.run()
    at.button(key=add).click()
    at.run()
    check(any("already assigned to another restatement group" in w.value for w in at.warning),
          "ui: a Claim already in a group is refused, with the reason")
    at.multiselect(key=pick).set_value([d, c])
    at.run()
    check(not any("already assigned" in w.value for w in at.warning),
          "ui: the refusal message goes once the selection changes")
    at.button(key=add).click()
    at.run()
    _flush(at)
    stored = store.load_review(review)
    groups = stored.active_restatement_groups()
    check(sorted(groups.values()) == sorted([sorted([a, b]), sorted([c, d])])
          and stored.response("SOURCE_RESTATEMENTS", source)["answer"] == "present",
          "ui: two groups are stored; the answer is present")
    check(book.count(sheets.RESTATEMENTS) == 4,
          "ui: nothing was written for the refused group")
    page = _markdown(at)
    check("**Group 1**" in page and "**Group 2**" in page
          and not any(w in page.lower() for w in ("canonical", "primary")),
          "ui: groups are listed with no canonical member")
    at.button(key="nav|review").click()
    at.run()
    page = _markdown(at)
    check("Restatements — 2 groups" in page, "ui: Review shows the number of groups")
    at.button(key="nav|recall").click()
    at.run()
    for key in [x.key for x in at.button if x.key and "|rg_rm|" in x.key]:
        at.button(key=key).click()
        at.run()
    _flush(at)
    check(store.load_review(review).active_restatement_groups() == {}
          and all(r["active"] == store.FALSE for r in book.read_tab(sheets.RESTATEMENTS)),
          "ui: removing groups marks their rows inactive")
    check(not at.radio(key=answer).disabled,
          "ui: No is available again once every group is removed")
    at.radio(key=answer).set_value("none")
    at.run()
    _flush(at)
    item = next(i for i in progress.recall_items(BRAIN, source, store.load_review(review))
                if i.question is spec.RESTATEMENTS)
    check(progress.item_problem(item, store.load_review(review)) is None,
          "ui: No after removing the groups completes Restatements")
    at.button(key="nav|review").click()
    at.run()
    check("Restatements — none" in _markdown(at), "ui: Review shows Restatements none")


def test_ui_calibration_and_doi():
    book = fresh_round()
    with_doi = next(s for s in _assigned(EVALUATOR, "training") if ui_doi(s))
    without = next((s for p in ("training", "agreement", "individual")
                    for s in _assigned(EVALUATOR, p) if not ui_doi(s)), None)
    phase = "training"
    review = rid(phase, EVALUATOR, with_doi)
    start(review)
    at = _app()
    _login(at)
    before = _stored_rows(book)
    _open(at, phase, with_doi)
    doi = str(BRAIN.source(with_doi)["doi"]).strip()
    line = [m.value for m in at.markdown if m.value.startswith("**DOI:**")]
    check(line and f"href='https://doi.org/{doi}'" in line[0]
          and "target='_blank'" in line[0] and "noopener" in line[0]
          and html.escape(doi) in line[0].split("<a")[0],
          "ui: the DOI stays visible, with an icon link to doi.org in a new tab")
    check(_stored_rows(book) == before, "ui: the DOI link writes nothing")
    import ui as ui_module
    target = "https://doi.org/10.1145/3696630.3728530"
    check(all(ui_module.doi_url(form) == target
              for form in ("10.1145/3696630.3728530", "https://doi.org/10.1145/3696630.3728530",
                           "http://doi.org/10.1145/3696630.3728530",
                           "doi:10.1145/3696630.3728530"))
          and ui_module.doi_url("unknown") == "" and ui_module.doi_url(None) == "",
          "ui: DOI forms normalise to https://doi.org/<doi>; none gives no link")
    if without:
        phase2 = next(p for p in ("training", "agreement", "individual")
                      if without in _assigned(EVALUATOR, p))
        start(rid(phase2, EVALUATOR, without))
        _open(at, phase2, without)
        line = [m.value for m in at.markdown if m.value.startswith("**DOI:**")]
        check(line and "<a" not in line[0], "ui: no DOI, no link icon")

    _open(at, phase, with_doi)
    at.button(key="nav|claims").click()
    at.run()
    instructions = [_inside(e) for e in _expanders(at, "Instructions")]
    general = [t for t in instructions if spec.CALIBRATION["claim.general"][0] in t]
    check(general and all(line in general[0] for line in spec.CALIBRATION["claim.general"])
          and "Schema / skill" not in general[0]
          and "This instruction applies" not in general[0],
          "ui: the Claim-level Instructions hold the three calibration lines only")
    check(all(not e.proto.expanded for e in _expanders(at, "Instructions")),
          "ui: calibration Instructions start closed")
    for field, key in (("basis", "claim.basis"),
                       ("claim_jurisdiction", "claim.claim_jurisdiction")):
        block = next((t for t in instructions
                      if html.escape(spec.plain(spec.definition(key)["text"])) in t), "")
        schema_at, calibration_at = block.find("Schema / skill"), block.find("Calibration")
        check(0 <= schema_at < calibration_at
              and all(html.escape(line) in block[calibration_at:]
                      for line in spec.CALIBRATION[field])
              and html.escape(spec.plain(spec.definition(key)["text"]))
              in block[schema_at:calibration_at],
              f"ui: {field} Instructions hold the schema text, then a separate Calibration")
    expected = {
        "claim.claim_object": (
            "Classify what the Claim asserts, not what it merely mentions.",
            "Law: the Claim is about law or legal/regulatory practice. This includes "
            "law regulating technology and the impact of technology on legal or "
            "regulatory practice.",
            "Technology: the Claim is about a technology, system or model. This "
            "includes how a technology behaves when applied to a legal task.",
            "Other: the Claim is substantively about both Law and Technology, or about "
            "Law or Technology together with another object that prevents either "
            "category from describing the Claim on its own."),
        "claim.basis": (
            "Abstract means abstract or conceptual considerations. It does not mean the "
            "Abstract section of the paper.",
            "Literature means that the Claim and its anchors present the point as "
            "coming from prior literature.",
            "Judge the Basis from the anchors attached to this Claim."),
        "claim.claim_jurisdiction": (
            "General: the Claim is explicitly jurisdiction-independent or concerns law "
            "in general.",
            "Undetermined: the jurisdiction is not stated and the context does not "
            "settle it.",
            "Do not choose General only because a technology Claim sounds broadly "
            "applicable."),
    }
    for key, lines in expected.items():
        block = next((t for t in instructions
                      if html.escape(spec.plain(spec.definition(key)["text"])) in t), "")
        calibration = block[block.find("Calibration"):]
        shown = re.findall(r"<p[^>]*>(.*?)</p>", calibration)
        check(shown == [html.escape(line) for line in lines],
              f"ui: {key} Calibration is exactly the agreed lines")
    page_text = _markdown(at)
    check("a general feature of the technology can be General" not in page_text
          and "presented as jurisdiction-independent" not in page_text,
          "ui: the previous technology-specific jurisdiction example is gone")
    claim_object = next((t for t in instructions
                         if html.escape(spec.plain(spec.definition("claim.claim_object")["text"]))
                         in t), "")
    schema_at, calibration_at = (claim_object.find("Schema / skill"),
                                 claim_object.find("Calibration"))
    frozen = claim_object[schema_at:calibration_at]
    check(0 <= schema_at < calibration_at
          and html.escape(spec.plain(spec.definition("claim.claim_object")["text"])) in frozen
          and all(html.escape(spec.plain(i["text"])) in frozen
                  for i in spec.definition("claim.claim_object")["items"])
          and all(v in frozen for v in ("Law", "Technology", "Other")),
          "ui: Claim object keeps its frozen definition, every category and meaning, "
          "under Schema / skill, before its Calibration")
    q6 = spec.BY_KEY["CLAIM_Q06_CLAIM_OBJECT"]
    check(q6.text == AGREED["CLAIM_Q06_CLAIM_OBJECT"] and q6.options == ("Yes", "No")
          and q6.field == "claim_object" and q6.calibration == ("claim_object",),
          "ui: the Claim object question, options and field are unchanged")
    q2 = next((t for t in instructions
               if html.escape(spec.plain(spec.definition("claim.node")["text"])) in t), "")
    check(q2 and "Calibration" not in q2, "ui: Question 2's Instructions are unchanged")


def ui_doi(source: str) -> bool:
    import ui as ui_module

    return bool(ui_module.doi_url(BRAIN.source(source).get("doi")))


# ====================================================================== main
def main() -> int:
    only = sys.argv[1] if len(sys.argv) > 1 else ""
    tests = [(name, fn) for name, fn in globals().items()
             if name.startswith("test_") and callable(fn) and only in name]
    for name, fn in tests:
        print(name)
        try:
            fn()
        except Exception:
            FAILURES.append(f"{name} raised")
            traceback.print_exc()
    print(f"\n{CHECKS} checks, {len(FAILURES)} failed")
    for failure in FAILURES:
        print(f"  - {failure}")
    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
